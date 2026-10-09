# -*- coding: utf-8 -*-
"""Genera el Manual de Usuario V3 del sistema de Transporte de Personal.

Cubre las novedades de la V3: prioridad de 3 niveles, emergencia medica
independiente, maestro de areas, filtros de fecha en Coordinacion y
Mis servicios, asignacion multiple visible, desasignar, auditoria en
Administracion y validaciones de cupo por fecha.
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MORADO = RGBColor(0x7B, 0x1F, 0xA2)
GRIS = RGBColor(0x55, 0x55, 0x55)
OUT = os.path.dirname(os.path.abspath(__file__))


def estilo_base(doc):
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    for h, sz in (('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12)):
        s = doc.styles[h]
        s.font.name = 'Calibri'
        s.font.size = Pt(sz)
        s.font.color.rgb = MORADO
        s.font.bold = True


def portada(doc, titulo, subtitulo):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo)
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = MORADO
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitulo)
    r2.font.size = Pt(14)
    r2.font.color.rgb = GRIS
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run('\nAplicativo de Transporte de Personal Interno\n'
               'Versión 3.0 — Octubre 2026').font.color.rgb = GRIS
    doc.add_page_break()


def tabla(doc, cabecera, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabecera):
        cell = t.rows[0].cells[i]
        cell.text = c
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
    for fila in filas:
        cells = t.add_row().cells
        for i, v in enumerate(fila):
            cells[i].text = str(v)
    if anchos:
        for i, w in enumerate(anchos):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def viñetas(doc, items):
    for it in items:
        doc.add_paragraph(it, style='List Bullet')


def pasos(doc, items):
    for it in items:
        doc.add_paragraph(it, style='List Number')


def indice(doc):
    """Inserta una tabla de contenido (campo TOC) con los títulos de nivel 1 y 2."""
    doc.add_heading('Índice', 1)
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    placeholder = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '(El índice se genera al abrir el documento en Word.)'
    placeholder.append(t)
    p._p.append(placeholder)
    run_end = p.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run_end._r.append(fld_end)
    doc.add_page_break()


def actualizar_campos_al_abrir(doc):
    """Hace que Word actualice los campos (índice) automáticamente al abrir."""
    upd = OxmlElement('w:updateFields')
    upd.set(qn('w:val'), 'true')
    doc.settings.element.append(upd)


def imagen(doc, archivo, caption, ancho=Cm(15.5)):
    ruta = os.path.join(OUT, 'img', 'v3', archivo)
    if not os.path.exists(ruta):
        ruta = os.path.join(OUT, 'img', 'v2', archivo)
    if not os.path.exists(ruta):
        ruta = os.path.join(OUT, 'img', archivo)
    if not os.path.exists(ruta):
        print('  (falta imagen: ' + archivo + ')')
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=ancho)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = GRIS


m = Document()
estilo_base(m)
portada(m, 'Manual de Usuario — V3', 'Aplicativo de Transporte de Personal Interno')
indice(m)

# ---------------- 1. INTRODUCCIÓN ----------------
m.add_heading('1. Introducción', 1)
m.add_paragraph(
    'Este manual explica paso a paso, con capturas de pantalla, cómo utilizar el aplicativo '
    'web de transporte de personal. El flujo es: un supervisor solicita una movilidad para su '
    'área, el jefe de transporte (coordinador) asigna las combis y las paradas, el conductor '
    'ejecuta el viaje y la administración consulta reportes y auditoría.')
m.add_paragraph('Requisitos para usar el aplicativo:')
viñetas(m, [
    'Un navegador moderno (Chrome, Edge o similar), en computadora o celular.',
    'Un usuario y contraseña asignados por el administrador del sistema.',
    'Conexión a internet para entrar; si la conexión se pierde mientras trabaja, la aplicación '
    'guarda sus acciones en el equipo y las envía automáticamente al volver (modo offline).',
    'Opcionalmente, aceptar el permiso de notificaciones del navegador para recibir avisos '
    'aunque esté en otra pestaña.',
])
m.add_paragraph(
    'Convenio del manual: cada figura muestra la pantalla real del sistema y debajo se indica su '
    'nombre ("Fig. N — descripción"). Los textos entre comillas corresponden a botones o '
    'etiquetas tal como aparecen en pantalla.')

m.add_heading('1.1 Novedades de la versión 3.0', 2)
viñetas(m, [
    'Prioridad de tres niveles: NORMAL, ALTA y EMERGENCIA. El supervisor la elige al crear la '
    'solicitud y define el orden de atención en Coordinación y Mis servicios.',
    'La marca "Emergencia médica" volvió a ser una casilla independiente de la prioridad: una '
    'solicitud puede ser urgente sin ser médica, o médica programada sin urgencia.',
    'Maestro de Áreas: las áreas se administran como catálogo (ya no texto libre). El supervisor '
    'solicita y ve únicamente las movilidades de su propia área.',
    'Filtro de fecha en Coordinación y Mis servicios, con botones "Todas" y "Hoy".',
    'Aviso de pendientes clickeable: muestra las fechas con solicitudes pendientes y salta al '
    'día más antiguo con un clic.',
    'Las filas pendientes se pintan en ámbar y las de emergencia en rojo para ubicarlas de '
    'inmediato.',
    'Cupo real por fecha: la disponibilidad de cada unidad se calcula sobre el día de la '
    'solicitud, incluyendo las porciones de asignaciones múltiples.',
    'Asignación múltiple visible: la solicitud repartida muestra la etiqueta MULTIPLE y las '
    'placas que la atienden.',
    'Nuevo botón "Desasignar" (↩): quita la unidad asignada y devuelve la solicitud a PENDIENTE.',
    'Banner "Ruta del día" del conductor: muestra la ruta real del traslado con sus paradas.',
    'Administración incorpora las pestañas Áreas y Auditoría (historial completo de acciones).',
])

# ---------------- 2. ACCESO ----------------
m.add_heading('2. Acceso al sistema', 1)
pasos(m, [
    'Abra el navegador e ingrese a la dirección del aplicativo (por ejemplo http://localhost:4200 '
    'o la URL interna que le indique el administrador).',
    'Escriba su Usuario y su Clave en el formulario de ingreso.',
    'Presione el botón "Ingresar".',
    'El sistema lo lleva automáticamente a su pantalla principal según su rol.',
])
imagen(m, 'v2_01_login.png',
       'Fig. 1 — Pantalla de inicio de sesión. El ojo junto al campo Clave permite ver lo que escribe.')
m.add_paragraph('La pantalla inicial que verá depende de su rol:')
tabla(m,
      ['Rol', 'Pantalla inicial', 'Opciones del menú'],
      [
          ['Supervisor', 'Solicitudes', 'Solicitudes (solo de su área)'],
          ['Jefe de Transporte', 'Coordinación', 'Coordinación, Reportes'],
          ['Conductor', 'Mis servicios', 'Mis servicios'],
          ['Administrador', 'Coordinación', 'Solicitudes, Coordinación, Mis servicios, Reportes, Administración'],
      ],
      anchos=[4.0, 4.0, 8.5])
m.add_paragraph(
    'Si el usuario o la clave son incorrectos aparece un aviso rojo bajo el formulario. Verifique '
    'los datos o pida al administrador que revise su cuenta en Administración → Usuarios.')
m.add_paragraph(
    'Para cerrar sesión presione "Cerrar sesión" (abajo del menú lateral) o el ícono de salida '
    'en el encabezado, junto a su nombre.')

# ---------------- 3. ELEMENTOS COMUNES ----------------
m.add_heading('3. Elementos comunes de la pantalla', 1)
m.add_paragraph(
    'Toda la aplicación comparte la misma estructura: a la izquierda el menú lateral con las '
    'opciones de su rol; arriba el encabezado con el nombre del módulo, la fecha, su nombre de '
    'usuario con su rol, y los íconos de notificaciones, cambio de clave y salir.')
tabla(m,
      ['Elemento', 'Dónde está', 'Para qué sirve'],
      [
          ['Menú lateral', 'Columna izquierda', 'Navegar entre los módulos de su rol. El botón ☰ lo pliega para ganar espacio.'],
          ['Campana', 'Encabezado, derecha', 'Abre el panel de notificaciones. Un punto rojo indica avisos sin leer; toque una para marcarla como leída.'],
          ['Badge amarillo', 'Encabezado, junto a la campana', 'Aparece solo si hay operaciones guardadas sin internet pendientes de sincronizar.'],
          ['Llave', 'Encabezado, junto a su nombre', 'Abre el formulario para cambiar su contraseña (mínimo 6 caracteres).'],
          ['Nombre + rol', 'Encabezado, derecha', 'Muestra su nombre y la etiqueta de su rol.'],
          ['Salir', 'Encabezado / menú lateral', 'Cierra la sesión y vuelve a la pantalla de ingreso.'],
      ],
      anchos=[3.2, 4.3, 9.0])
imagen(m, 'v2_05_notificaciones.png',
       'Fig. 2 — Panel de notificaciones al presionar la campana.')
m.add_paragraph('El sistema le avisa cuando ocurren estos eventos:')
viñetas(m, [
    'Le asignan una unidad a su solicitud (al supervisor que la creó y al conductor de la placa).',
    'Su solicitud se une o se acopla a un traslado.',
    'Su solicitud es separada de un traslado o desasignada de su unidad.',
    'El estado de su solicitud cambia (en ruta, realizada, etc.).',
    'El conductor agrega pasajeros de emergencia a un servicio.',
])

m.add_heading('3.1 Las tablas y su paginador', 2)
viñetas(m, [
    '"Mostrando 1 a 10 de N registros": indica qué filas está viendo y el total.',
    'Flechas « ‹ › »: cambiar de página (primera, anterior, siguiente, última).',
    'Selector de filas: elija ver 5, 10, 20 o 50 registros por página.',
    'Encabezados en mayúsculas; la última columna es siempre ACCIONES cuando la tabla tiene botones.',
])

m.add_heading('3.2 Botones de acción con íconos', 2)
m.add_paragraph(
    'Los botones de acción de las tablas son solo íconos, todos del mismo tamaño. Para saber qué '
    'hace cada uno:')
viñetas(m, [
    'En computadora: pase el mouse sobre el botón y aparece una etiqueta flotante (tooltip) '
    'con el nombre de la acción.',
    'En celular: el nombre de la acción se muestra siempre debajo del ícono.',
])
tabla(m,
      ['Ícono', 'Acción', 'Dónde aparece'],
      [
          ['▶ (azul)', 'Iniciar el servicio', 'Mis servicios'],
          ['✓ (verde)', 'Finalizar el servicio', 'Mis servicios'],
          ['Persona+ (rojo)', 'Agregar pasajeros de emergencia', 'Mis servicios'],
          ['Mapa (celeste)', 'Ver ruta y paradas', 'Mis servicios'],
          ['Lápiz (azul)', 'Editar el registro', 'Solicitudes, Coordinación, Administración'],
          ['X (roja)', 'Anular o desactivar', 'Solicitudes, Coordinación'],
          ['⛓ (roja)', 'Separar del traslado', 'Coordinación'],
          ['⚡ (verde)', 'Asignación automática de unidad', 'Coordinación'],
          ['⧉ (gris)', 'Asignación múltiple de unidades', 'Coordinación'],
          ['↩ (ámbar)', 'Desasignar la(s) unidad(es)', 'Coordinación'],
      ],
      anchos=[3.2, 6.5, 6.0])

m.add_heading('3.3 Colores de las filas y etiquetas', 2)
tabla(m,
      ['Color / etiqueta', 'Significado'],
      [
          ['Fila ámbar + borde naranja', 'Solicitud PENDIENTE: requiere atención del coordinador.'],
          ['Fila roja', 'Solicitud con prioridad EMERGENCIA.'],
          ['Fila verde', 'Solicitud agrupada dentro de un traslado (T-N).'],
          ['Etiqueta ALTA (ámbar)', 'Prioridad alta: se atiende antes que las normales.'],
          ['Etiqueta MÉDICA (roja, corazón)', 'La solicitud es una emergencia médica.'],
          ['Etiqueta MULTIPLE (ámbar)', 'La solicitud está repartida entre varias unidades; las placas se ven al lado o con el tooltip.'],
      ],
      anchos=[5.5, 11.0])

m.add_heading('3.4 Uso en celular', 2)
viñetas(m, [
    'El menú lateral se repliega y se abre con el botón ☰ del encabezado.',
    'Los paneles de notificaciones y de cambio de clave ocupan casi todo el ancho.',
    'Las tablas anchas se recorren deslizando el dedo hacia los lados.',
    'En Administración, cada fila se convierte en una tarjeta con sus datos etiquetados.',
    'Los botones de acción muestran su nombre debajo del ícono.',
])
imagen(m, 'v2_m01_solicitudes.png', 'Fig. 3 — Solicitudes en celular: filtros, tabla y paginador.', Cm(7.0))
imagen(m, 'v2_m04_administracion.png',
       'Fig. 4 — Administración en celular: cada registro se muestra como tarjeta etiquetada.', Cm(7.0))

m.add_heading('3.5 Trabajar sin internet (modo offline)', 2)
m.add_paragraph(
    'Si pierde la conexión, la aplicación sigue funcionando con la última información cargada: '
    'puede consultar listados que ya haya visto y ejecutar acciones (crear, asignar, cambiar '
    'estado, etc.). Estas acciones se guardan en el equipo y aparecen contadas en el badge '
    'amarillo del encabezado. Al recuperar internet se envían solas en el mismo orden y el badge '
    'desaparece cuando todo quedó sincronizado. Si una acción ya no es válida al sincronizar '
    '(por ejemplo, el cupo se llenó mientras tanto), el sistema la descarta de la cola.')

# ---------------- 4. SOLICITUDES ----------------
m.add_heading('4. Módulo Solicitudes (Supervisor)', 1)
m.add_paragraph(
    'Es la pantalla principal del supervisor. Muestra únicamente las solicitudes de su área '
    '(el sistema la toma automáticamente de su usuario) y permite crear, editar, anular y '
    'acoplar solicitudes.')
imagen(m, 'v3_solicitudes.png',
       'Fig. 5 — Listado de solicitudes del supervisor: solo su área, columna PRIORIDAD, '
       'CUPO de la unidad asignada y etiqueta MULTIPLE cuando se repartió entre varias.')

m.add_heading('4.1 Entender el listado', 2)
tabla(m,
      ['Columna', 'Qué muestra'],
      [
          ['PRIORIDAD', 'NORMAL, ALTA o EMERGENCIA, más la etiqueta MÉDICA si aplica. Emergencias y altas van primero.'],
          ['NOMBRE', 'Quién registró la solicitud.'],
          ['FECHA / HORA', 'Fecha y hora programada del viaje.'],
          ['ÁREA', 'Área del solicitante (siempre la suya).'],
          ['P. PARTIDA / P. LLEGADA', 'Punto de origen y punto de destino (del catálogo de Puntos).'],
          ['CANTIDAD', 'Número de personas a trasladar.'],
          ['MOTIVO', 'Motivo del viaje (capacitación, actividad, salud, etc.).'],
          ['PLACA', 'Unidad asignada, o MULTIPLE si se repartió entre varias.'],
          ['CUPO', 'Ocupadas/capacidad de la unidad. Verde = hay espacio, ámbar = casi llena, rojo = llena.'],
          ['ESTADO', 'PENDIENTE, ASIGNADO, EN_RUTA, REALIZADO o ANULADO.'],
          ['ACCIONES', 'Lápiz = editar (solo pendientes). X roja = anular (pide confirmación).'],
      ],
      anchos=[4.0, 12.5])
m.add_paragraph(
    'Sobre la tabla hay dos campos de fecha "Desde" / "Hasta": al elegir un rango la lista se '
    'limita a esas fechas, y aparece el botón "Ver todas" para quitar el filtro. Clic sobre '
    'cualquier fila abre el detalle de la solicitud.')

m.add_heading('4.2 Crear una solicitud', 2)
pasos(m, [
    'Presione el botón "Crear" en la parte superior derecha.',
    'Su nombre y su área se cargan automáticamente (no se pueden modificar).',
    'Complete los campos obligatorios (marcados con *):',
])
tabla(m,
      ['Campo', 'Obligatorio', 'Descripción'],
      [
          ['FECHA PROGRAMADA', 'Sí', 'Día en que se realizará el viaje.'],
          ['Hr. PROGRAMADA', 'Sí', 'Hora de salida.'],
          ['P. PARTIDA', 'Sí', 'Punto donde sube el personal (lista del catálogo de Puntos).'],
          ['P. LLEGADA', 'Sí', 'Punto de destino.'],
          ['CANTIDAD', 'Sí', 'Personas a trasladar (mínimo 1; con flechas para subir/bajar).'],
          ['MOTIVO', 'Sí', 'Motivo del viaje (lista del catálogo de Motivos).'],
          ['OBSERVACIÓN', 'No', 'Comentario libre para el coordinador o el conductor.'],
          ['PRIORIDAD', 'Sí', 'Normal (por defecto), Alta o Emergencia. Define el orden de atención.'],
          ['Emergencia médica', 'No', 'Casilla independiente: indica que el traslado es por un caso médico. Al marcarla sugiere prioridad EMERGENCIA, editable.'],
      ],
      anchos=[4.5, 2.5, 9.5])
m.add_paragraph(
    'La diferencia entre ambos campos: la PRIORIDAD indica qué tan urgente es atender la '
    'solicitud; "Emergencia médica" indica de qué tipo de caso se trata (dato que el conductor '
    'y los reportes necesitan ver). Una cita médica programada puede ser médica con prioridad '
    'NORMAL; una auditoría urgente puede ser prioridad ALTA sin ser médica.')
pasos(m, [
    'Presione "Save". El botón "Cancel" cierra sin guardar.',
    'La solicitud aparece en el listado con estado PENDIENTE; el coordinador ya puede verla.',
])
imagen(m, 'v3_crear_solicitud.png',
       'Fig. 6 — Formulario de nueva solicitud: combo PRIORIDAD junto a la casilla '
       '"Emergencia médica".')

m.add_heading('4.3 Ver el detalle de una solicitud', 2)
m.add_paragraph(
    'Haga clic sobre la fila de la solicitud. Se abre una ventana con todos sus datos: prioridad, '
    'solicitante, fecha y hora, área, puntos, cantidad, motivo, observación, placa, estado y si ya '
    'fue realizada. Si la solicitud está unida a un traslado, la ventana muestra además la ruta, '
    'las paradas con cuántos suben y bajan en cada una, y las demás solicitudes que comparten el viaje.')
imagen(m, 'v2_04_detalle_solicitud.png', 'Fig. 7 — Detalle completo de una solicitud (clic sobre la fila para abrirlo).')

m.add_heading('4.4 Acoplarse a una movilidad existente', 2)
m.add_paragraph(
    'Si ya existe un traslado programado que pasa cerca de su punto de partida y destino, puede '
    'unir su solicitud a ese viaje sin esperar una nueva asignación. Solo puede acoplar '
    'solicitudes de su propia área.')
pasos(m, [
    'Cree su solicitud y ábrala haciendo clic en su fila.',
    'Al final del detalle aparece la sección "Acoplar a traslado con cupo" con los traslados '
    'compatibles (origen y destino a menos de 10 km de la ruta).',
    'Cada botón muestra el traslado (T-N), la placa y su ruta. Presione el que le sirva.',
    'El sistema verifica que haya cupo real para esa fecha; si lo hay, su solicitud queda '
    'ASIGNADA a ese traslado conservando su registro individual.',
])

m.add_heading('4.5 Editar o anular una solicitud', 2)
viñetas(m, [
    'Editar (lápiz azul): abre el formulario con los datos cargados para corregirlos. Disponible '
    'solo mientras la solicitud está PENDIENTE; una vez asignada, solo el coordinador puede '
    'modificarla o desasignarla.',
    'Anular (X roja): el sistema pide confirmación antes de cancelar. La acción no se puede '
    'deshacer y queda registrada en la auditoría. Las solicitudes realizadas no se pueden anular.',
])

m.add_heading('4.6 Estados de una solicitud', 2)
tabla(m,
      ['Estado', 'Qué significa', 'Quién lo produce'],
      [
          ['PENDIENTE', 'Recién creada, aún sin unidad asignada.', 'Supervisor'],
          ['ASIGNADO', 'Ya tiene placa (sola, dentro de un traslado o repartida en varias unidades).', 'Coordinador / acople'],
          ['EN_RUTA', 'El conductor inició el recorrido.', 'Conductor'],
          ['REALIZADO', 'El conductor finalizó el viaje; se guardan las horas.', 'Conductor'],
          ['ANULADO', 'Solicitud cancelada.', 'Supervisor / coordinador'],
      ],
      anchos=[3.0, 9.0, 4.5])

# ---------------- 5. COORDINACIÓN ----------------
m.add_heading('5. Módulo Coordinación (Jefe de Transporte)', 1)
m.add_paragraph(
    'Es el centro de operaciones del día. La pantalla se divide en dos paneles: a la izquierda '
    'las solicitudes con sus acciones, y a la derecha los traslados activos con su ocupación.')
imagen(m, 'v3_coordinacion.png',
       'Fig. 8 — Coordinación al día de hoy: selector de fecha junto al título, aviso de '
       'pendientes, badges de cupo, columna FECHA y la solicitud repartida con etiqueta '
       'MULTIPLE y sus dos placas.')

m.add_heading('5.1 Filtro de fecha y aviso de pendientes', 2)
m.add_paragraph(
    'Junto al título "Solicitudes" hay un selector de fecha (por defecto el día de hoy). Todo lo '
    'que se ve en pantalla corresponde a esa fecha: la lista de solicitudes, los traslados '
    'activos y los badges de cupo de las unidades.')
viñetas(m, [
    'Cambie la fecha para gestionar otro día (por ejemplo, preparar las movilidades de mañana).',
    'El botón "Todas" muestra todas las solicitudes de todas las fechas; cambia a "Hoy" para '
    'volver. En modo "Todas" los badges de cupo se ocultan porque dependen de una fecha.',
    'Si hay solicitudes PENDIENTE en otras fechas, aparece un aviso ámbar "Pendientes: '
    'dd/MM (N)" junto al selector. Presiónelo para saltar directo a la fecha pendiente más '
    'antigua.',
    'Las filas pendientes se pintan en ámbar para distinguirlas de un vistazo.',
])
m.add_paragraph(
    'La tabla muestra una columna por cada dato de la solicitud: NOMBRE, FECHA (día/mes), '
    'Hr., P. PARTIDA, P. LLEGADA, CANT., MOTIVO (con las etiquetas de prioridad y MÉDICA), '
    'PLACA (o la etiqueta MULTIPLE con las placas en TRASLADO), TRASLADO (T-N si está unida) '
    'y ACCIONES.')
imagen(m, 'v3_coordinacion_fecha.png',
       'Fig. 9 — Coordinación viendo otra fecha (05/10): la fila pendiente se pinta en ámbar, '
       'los badges muestran el cupo de ese día y la solicitud ASIGNADO ofrece el botón '
       'Desasignar (↩).')

m.add_heading('5.2 Ver la disponibilidad de las unidades', 2)
m.add_paragraph(
    'Sobre la tabla hay un badge por cada combi con la forma "PLACA · libres/capacidad", por '
    'ejemplo "TP-001 · 15/15 libres". El cupo corresponde a la fecha seleccionada e incluye lo '
    'ocupado por asignaciones simples y por porciones de asignaciones múltiples. El color '
    'resume el estado:')
viñetas(m, [
    'Verde: la unidad está libre o con bastante cupo.',
    'Ámbar: parcialmente ocupada.',
    'Rojo: llena — no acepta más pasajeros.',
])
m.add_paragraph(
    'Los selectores de placa de cada fila muestran el cupo en la fecha de esa solicitud (con la '
    'fecha entre paréntesis cuando no es la seleccionada) y bloquean las unidades sin espacio.')

m.add_heading('5.3 Asignar unidad a una solicitud', 2)
tabla(m,
      ['Botón', 'Acción'],
      [
          ['Selector "Placa" + ✓ azul', 'Elige la unidad en la lista (muestra el cupo de cada una) y presione el check para asignarla.'],
          ['⚡ (rayo verde)', 'Asignación automática: toma la primera unidad con cupo suficiente para esa fecha.'],
          ['⧉ (ventanas gris)', 'Asignación múltiple: reparte la cantidad entre varias unidades cuando una sola no alcanza.'],
          ['Lápiz (azul)', 'Editar los datos de la solicitud (solo pendientes).'],
          ['X (roja)', 'Anular la solicitud.'],
          ['↩ (ámbar)', 'Desasignar: quita la(s) unidad(es) de una solicitud ASIGNADA y la devuelve a PENDIENTE.'],
          ['⛓ (roja)', 'Separar del traslado (solo en solicitudes ya unidas y aún no iniciadas).'],
      ],
      anchos=[5.0, 11.5])
pasos(m, [
    'En la columna ACCIONES de la solicitud elija la placa en el selector.',
    'Presione el botón ✓. Confirme el mensaje.',
    'La solicitud pasa a ASIGNADO y el conductor de esa placa recibe una notificación.',
    'Para corregir una asignación use ↩ (Desasignar) y vuelva a asignar.',
])
m.add_paragraph(
    'Asignación múltiple (⧉): para solicitudes que superan la capacidad de una combi. Se abre un '
    'panel donde agrega tantas unidades como necesite con la cantidad que llevará cada una; el '
    'total repartido debe coincidir exactamente con lo solicitado y cada unidad valida su cupo '
    'libre real para la fecha. En la tabla la solicitud muestra la etiqueta MULTIPLE y las '
    'placas asignadas en la columna TRASLADO (o con el tooltip).')
imagen(m, 'v3_asignacion_multiple.png',
       'Fig. 10 — Diálogo "Distribuir N personas entre varias unidades": una fila por unidad '
       'con su tope de cupo y el contador de personas asignadas.')

m.add_heading('5.4 Unir solicitudes en un traslado', 2)
pasos(m, [
    'Marque con los checkboxes de la primera columna dos o más solicitudes pendientes de la '
    'misma fecha (solo las pendientes pueden seleccionarse).',
    'Arriba de la tabla aparecen el total de personas, el selector "Placa del traslado" y el '
    'botón "Unir (N)".',
    'Elija la placa y presione "Unir". Se crea el traslado T-N que agrupa las solicitudes: cada '
    'una conserva su número, solicitante y cantidad para los reportes.',
])
m.add_paragraph(
    'Separar: en la fila de una solicitud que está dentro de un traslado aparece el botón de '
    'desenlazar (⛓). Tras confirmar, la solicitud vuelve a estado PENDIENTE sin placa y el '
    'traslado recalcula su ocupación. Si el traslado queda vacío, se anula solo.')

m.add_heading('5.5 Definir las paradas y la ruta', 2)
pasos(m, [
    'En la tarjeta del traslado (panel derecho) presione "Editar paradas".',
    'Con "Agregar parada" añada filas: en cada una elija el punto (lista con buscador) y escriba '
    'cuántas personas suben y cuántas bajan ahí.',
    'Use las flechas ↑ ↓ para ordenar la secuencia de recogida y el basurero para quitar una parada.',
    'Presione "Guardar ruta". El conductor verá las paradas en ese orden en su aplicación y en '
    'su banner "Ruta del día".',
])
m.add_paragraph(
    'Con "Mapa" se despliega el recorrido sobre OpenStreetMap; "OSM" abre la ruta en una pestaña '
    'nueva para compartirla o navegarla.')

m.add_heading('5.6 Editar solicitudes', 2)
m.add_paragraph(
    'El coordinador puede modificar cualquier solicitud pendiente. El lápiz abre un panel con '
    'todos los campos: fecha, hora, puntos, cantidad, motivo, observación, PRIORIDAD y la marca '
    'de emergencia médica. Todo cambio queda en auditoría.')

# ---------------- 6. MIS SERVICIOS ----------------
m.add_heading('6. Módulo Mis servicios (Conductor)', 1)
m.add_paragraph(
    'El conductor solo ve los servicios asignados a la placa de su unidad, ordenados con las '
    'emergencias primero. Arriba hay un selector de fecha (por defecto hoy) con botones '
    '"Todas" / "Hoy", el botón "Actualizar" y la franja "Ruta del día" con la secuencia de '
    'puntos del día seleccionado (usa la ruta real del traslado cuando existe). En modo "Todas" '
    'el banner se oculta.')
imagen(m, 'v3_conductor.png',
       'Fig. 11 — Mis servicios: selector de fecha con botón "Todas", banner "Ruta del día" '
       'y la porción de 5 personas que le corresponde a esta combi de una solicitud repartida.')
tabla(m,
      ['Botón (columna ACCIONES)', 'Qué hace'],
      [
          ['▶ Iniciar (azul)', 'Aparece en servicios ASIGNADO. Pasa a EN_RUTA y guarda la hora de inicio.'],
          ['✓ Finalizar (verde)', 'Aparece en servicios EN_RUTA. Pasa a REALIZADO y guarda la hora de fin.'],
          ['Persona+ (rojo)', 'Agregar pasajeros de emergencia si hay cupo libre.'],
          ['Mapa (celeste)', 'Ver la ruta del servicio con sus paradas ordenadas.'],
      ],
      anchos=[5.0, 11.5])
pasos(m, [
    'Al llegar a su turno, abra "Mis servicios" y revise los viajes del día: origen, destino, '
    'cantidad de personas, motivo y si es emergencia médica.',
    'Si un servicio es parte de una solicitud repartida en varias unidades, verá su porción con '
    'la cantidad de personas que le corresponde a su combi.',
    'Antes de arrancar, presione el ícono del mapa para conocer la ruta completa.',
    'Al iniciar el recorrido presione ▶ (Iniciar). El estado cambia a EN_RUTA.',
    'Recorra las paradas en el orden indicado.',
    'Al terminar el viaje presione ✓ (Finalizar). El servicio queda REALIZADO. En solicitudes '
    'repartidas entre varias combis, cada conductor finaliza su porción; la solicitud se marca '
    'REALIZADO cuando todas las porciones terminaron.',
    'Sin señal puede seguir operando: los cambios se guardan en el equipo y se envían al '
    'recuperar internet.',
])
imagen(m, 'v2_08_ruta_paradas.png',
       'Fig. 12 — Ruta del servicio: mapa con la trayectoria y tabla de paradas ordenadas.')
m.add_heading('6.1 Agregar pasajeros de emergencia', 2)
pasos(m, [
    'En un servicio ASIGNADO o EN_RUTA presione el botón rojo de persona+.',
    'El panel muestra el cupo libre de la unidad.',
    'Escriba cuántos pasajeros adicionales suben y presione "Agregar".',
    'Si la cantidad supera el cupo disponible el sistema lo rechaza. Al guardar, esa porción '
    'del servicio queda marcada como emergencia y el supervisor recibe una notificación.',
])

# ---------------- 7. REPORTES ----------------
m.add_heading('7. Módulo Reportes (Jefe / Administrador)', 1)
m.add_paragraph('Genera el histórico de movilidades y los indicadores de operación.')
pasos(m, [
    'Elija el rango "Desde" / "Hasta" y presione "Generar".',
    'Opcionalmente refine con los filtros Supervisor, Placa y Estado.',
    'Revise las tarjetas de indicadores y las tablas de eficiencia.',
    'Use "Exportar CSV" para descargar el detalle (compatible con Excel) o "Imprimir / PDF" '
    'para la versión imprimible.',
])
imagen(m, 'v3_reportes.png',
       'Fig. 13 — Reportes: filtros de fecha y supervisor/placa/estado, totales rápidos y '
       'tarjetas de indicadores de operación.')
tabla(m,
      ['Indicador', 'Qué mide'],
      [
          ['Solicitudes / Personas', 'Total de movilidades y de personas trasladadas en el rango.'],
          ['Realizados / Emergencias', 'Viajes completados y cuántos fueron de emergencia.'],
          ['Duración promedio', 'Tiempo medio entre Iniciar y Finalizar.'],
          ['Espera promedio', 'Tiempo medio entre la creación de la solicitud y su atención.'],
          ['Puntualidad ±10 min', 'Servicios realizados dentro de 10 minutos de la hora programada.'],
          ['Eficiencia por unidad', 'Viajes, personas y ocupación por combi (incluye las porciones de asignaciones múltiples).'],
          ['Servicios por conductor', 'Viajes y personas atendidas por cada conductor.'],
          ['Movimiento por área', 'Solicitudes y personas agrupadas por el área del maestro.'],
      ],
      anchos=[5.0, 11.5])
m.add_paragraph(
    'La columna de prioridad del detalle muestra NORMAL / ALTA / EMERGENCIA y la etiqueta MÉDICA '
    'cuando corresponde; ambas se exportan al CSV.')
m.add_heading('7.1 Auditoría', 2)
m.add_paragraph(
    'El botón "Cargar auditoría" muestra el historial de todas las acciones del sistema: inicios '
    'de sesión, creación y edición de solicitudes, asignaciones, uniones, separaciones, cambios '
    'de estado y pasajeros agregados, con fecha, entidad afectada y usuario. El mismo historial '
    'está disponible para el administrador en la pestaña Auditoría de Administración.')
m.add_paragraph(
    'La columna Detalle guarda el JSON exacto de la operación (qué se envió al sistema), útil '
    'para rastrear por ejemplo qué placas y cantidades se repartieron en una asignación múltiple.')

# ---------------- 8. ADMINISTRACIÓN ----------------
m.add_heading('8. Módulo Administración (Administrador)', 1)
m.add_paragraph(
    'Gestiona los catálogos y la auditoría del sistema en siete pestañas. En las seis primeras '
    'funciona igual: el formulario de arriba sirve para crear (botón "Crear") y para editar: al '
    'presionar el lápiz en una fila, sus datos se cargan en el formulario y el botón cambia a '
    '"Actualizar". El botón de prohibido inhabilita el registro sin borrarlo; en usuarios, si '
    'deja la clave vacía al editar se conserva la actual.')
tabla(m,
      ['Pestaña', 'Campos / contenido', 'Para qué sirve'],
      [
          ['Usuarios', 'Usuario, nombre, rol, área (del maestro) y clave.',
           'Cuentas de supervisores, coordinadores y administradores.'],
          ['Conductores', 'Usuario, nombre, placa de la unidad asignada, área y clave.',
           'Cuentas de conductor; la placa define qué servicios verá en "Mis servicios".'],
          ['Unidades', 'Placa y capacidad.',
           'Las combis disponibles para asignar.'],
          ['Puntos', 'Nombre, latitud y longitud.',
           'Orígenes/destinos y paradas. Las coordenadas alimentan el mapa y la validación de rutas al acoplarse.'],
          ['Motivos', 'Nombre.',
           'Catálogo de motivos de viaje que elige el supervisor.'],
          ['Áreas', 'Nombre.',
           'Maestro de áreas: se asigna a usuarios y conductores; define qué solicitudes ve cada supervisor.'],
          ['Auditoría', 'Fecha, entidad, acción, usuario y detalle (solo lectura).',
           'Historial completo de todas las acciones del sistema.'],
      ],
      anchos=[2.8, 6.7, 7.0])
imagen(m, 'v3_administracion.png',
       'Fig. 14 — Administración: siete pestañas, formulario arriba (con el combo Área del '
       'maestro) y tabla con paginador.')
imagen(m, 'v3_areas.png',
       'Fig. 15 — Pestaña Áreas: catálogo maestro que se asigna a usuarios y conductores.')
imagen(m, 'v3_auditoria.png',
       'Fig. 16 — Pestaña Auditoría: historial de acciones con fecha, entidad, usuario y '
       'detalle de la operación.')
m.add_paragraph(
    'Importante sobre las áreas: asígnele un área a cada supervisor desde el combo (no es texto '
    'libre). Un supervisor sin área solo verá las solicitudes que él mismo registró. El área de '
    'la solicitud se toma automáticamente del supervisor que la crea.')

# ---------------- 9. GUÍA RÁPIDA ----------------
m.add_heading('9. Guía rápida por rol', 1)
tabla(m,
      ['Rol', 'Flujo típico'],
      [
          ['Supervisor',
           '1. Ingresar → Solicitudes (solo su área). 2. "Crear", elegir prioridad y marcar '
           'emergencia médica si aplica. 3. (Opcional) Abrir el detalle y acoplarse a un '
           'traslado con cupo. 4. Seguir el ESTADO. 5. Editar o anular mientras esté pendiente.'],
          ['Jefe de Transporte',
           '1. Ingresar → Coordinación. 2. Elegir la fecha y atender el aviso de pendientes. '
           '3. Asignar placa (✓, ⚡ o ⧉) o unir solicitudes compatibles en un traslado. '
           '4. "Editar paradas" para ordenar el recorrido. 5. Desasignar (↩) para corregir. '
           '6. Revisar Reportes al cierre.'],
          ['Conductor',
           '1. Ingresar → Mis servicios (fecha de hoy o "Todas"). 2. Revisar la "Ruta del día" '
           'y las paradas con el ícono de mapa. 3. ▶ al arrancar. 4. Persona+ solo para '
           'pasajeros de emergencia con cupo. 5. ✓ al terminar su porción del viaje.'],
          ['Administrador',
           '1. Mantener catálogos: usuarios con su área, conductores con su placa, unidades, '
           'puntos con coordenadas, motivos y áreas. 2. Revisar Reportes y Auditoría. '
           '3. Puede operar cualquier pantalla de los demás roles.'],
      ],
      anchos=[3.5, 13.0])

# ---------------- 10. GLOSARIO ----------------
m.add_heading('10. Glosario', 1)
tabla(m,
      ['Término', 'Significado'],
      [
          ['Solicitud', 'Pedido de transporte creado por un supervisor para su área.'],
          ['Traslado (T-N)', 'Agrupación de una o más solicitudes bajo una misma placa, con su ruta y paradas ordenadas.'],
          ['Acoplarse', 'Unir una solicitud pendiente a un traslado con cupo y ruta compatible (máx. 10 km).'],
          ['Asignación múltiple', 'Reparto de una solicitud entre varias unidades; la etiqueta MULTIPLE lo indica.'],
          ['Porción', 'Parte de una solicitud múltiple que atiende una unidad concreta.'],
          ['Parada', 'Punto intermedio de la ruta donde sube o baja personal.'],
          ['Cupo', 'Asientos libres de la unidad en una fecha (ocupadas/capacidad).'],
          ['Prioridad', 'Nivel de urgencia de la solicitud: NORMAL, ALTA o EMERGENCIA.'],
          ['Emergencia médica', 'Marca que indica que el traslado responde a un caso médico; independiente de la prioridad.'],
          ['Área', 'Unidad organizativa del supervisor; define qué solicitudes ve y cómo se agrupan los reportes.'],
          ['Badge', 'Etiqueta de color que resume un estado (verde/ámbar/rojo).'],
          ['Tooltip', 'Etiqueta flotante con el nombre de un botón; aparece al pasar el mouse.'],
          ['Paginador', 'Controles al pie de cada tabla para cambiar de página y elegir cuántas filas ver.'],
          ['Auditoría', 'Registro de quién hizo cada cambio y cuándo, con el detalle de la operación.'],
          ['Modo offline', 'Trabajo sin internet: las acciones se guardan en el equipo y se sincronizan al volver la conexión.'],
      ],
      anchos=[3.5, 13.0])

actualizar_campos_al_abrir(m)
m.save(os.path.join(OUT, 'Manual_de_Usuario_Transporte_Personal_V3.docx'))
print('OK Manual V3')

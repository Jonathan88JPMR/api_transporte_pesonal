# -*- coding: utf-8 -*-
"""Genera los manuales V4 del aplicativo Transporte de Personal:

  1. Manual_General_Transporte_Personal_V4.docx  -> todo el sistema
  2. Manual_Rol_Supervisor.docx                   -> modulo Solicitudes
  3. Manual_Rol_Coordinador.docx                  -> Coordinacion + Dashboard + Reportes + Administracion (limitada)
  4. Manual_Rol_Conductor.docx                    -> Mis servicios
  5. Manual_Rol_Administrador.docx                -> todo

Las capturas se toman de docs/img/v4 (Playwright, Oct-2026).
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MORADO = RGBColor(0x7B, 0x1F, 0xA2)
VERDE = RGBColor(0x2E, 0x7D, 0x32)
GRIS = RGBColor(0x55, 0x55, 0x55)
OUT = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(OUT, 'img', 'v4')

VERSION = 'Versión 4.0 — Octubre 2026'


# ===================== utilidades =====================

def estilo_base(doc):
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    for h, sz in (('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12)):
        s = doc.styles[h]
        s.font.name = 'Calibri'
        s.font.size = Pt(sz)
        s.font.color.rgb = MORADO
        s.font.bold = True


def portada(doc, titulo, subtitulo):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo)
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = VERDE
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitulo)
    r2.font.size = Pt(14)
    r2.font.color.rgb = GRIS
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run('\nAplicativo de Transporte de Personal Interno\n' + VERSION).font.color.rgb = GRIS
    doc.add_page_break()


def tabla(doc, cabecera, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabecera):
        cell = t.rows[0].cells[i]
        cell.text = c
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
    for fila in filas:
        cells = t.add_row().cells
        for i, v in enumerate(fila):
            cells[i].text = str(v)
    if anchos:
        for i, w in enumerate(anchos):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def viñetas(doc, items):
    for it in items:
        doc.add_paragraph(it, style='List Bullet')


def pasos(doc, items):
    for it in items:
        doc.add_paragraph(it, style='List Number')


def indice(doc):
    doc.add_heading('Índice', 1)
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    placeholder = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '(El índice se genera al abrir el documento en Word: clic derecho > Actualizar campos.)'
    placeholder.append(t)
    p._p.append(placeholder)
    run_end = p.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run_end._r.append(fld_end)
    doc.add_page_break()


def actualizar_campos_al_abrir(doc):
    upd = OxmlElement('w:updateFields')
    upd.set(qn('w:val'), 'true')
    doc.settings.element.append(upd)


class Ctx:
    """Envuelve el documento y lleva el contador de figuras."""
    def __init__(self):
        self.doc = Document()
        estilo_base(self.doc)
        self.nfig = 0

    def imagen(self, archivo, caption, ancho=Cm(15.5)):
        ruta = os.path.join(IMG, archivo)
        if not os.path.exists(ruta):
            ruta = os.path.join(OUT, 'img', archivo)
        if not os.path.exists(ruta):
            print('  (falta imagen: ' + archivo + ')')
            return
        self.nfig += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(ruta, width=ancho)
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run('Fig. %d — %s' % (self.nfig, caption))
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = GRIS

    def save(self, nombre):
        actualizar_campos_al_abrir(self.doc)
        ruta = os.path.join(OUT, nombre)
        self.doc.save(ruta)
        print('OK ' + nombre)


# ===================== secciones reutilizables =====================

def s_intro(c, alcance):
    d = c.doc
    d.add_heading('Introducción', 1)
    d.add_paragraph(
        'El aplicativo web de Transporte de Personal organiza las movilidades internas del '
        'personal. El flujo general es: un supervisor solicita una movilidad para su área, el '
        'coordinador (jefe de transporte) asigna las combis y define las paradas, el conductor '
        'ejecuta el viaje marcando cada parada y la administración consulta reportes, dashboard '
        'y auditoría.')
    d.add_paragraph(alcance)
    d.add_paragraph('Requisitos para usar el aplicativo:')
    viñetas(d, [
        'Un navegador moderno (Chrome, Edge o similar), en computadora o celular.',
        'Un usuario y contraseña asignados por el administrador del sistema.',
        'Conexión a internet para ingresar. Si la conexión se pierde mientras trabaja, la '
        'aplicación guarda sus acciones en el equipo y las envía automáticamente al volver '
        '(modo offline).',
        'Opcionalmente, aceptar el permiso de notificaciones del navegador para recibir avisos '
        'aunque esté en otra pestaña.',
        'En el caso del conductor, aceptar el permiso de ubicación (GPS): la aplicación reporta '
        'la posición de la unidad cada 30 segundos mientras un servicio está en ruta.',
    ])
    d.add_paragraph(
        'Convenio del manual: cada figura muestra la pantalla real del sistema con el número de '
        'figura debajo ("Fig. N — descripción"). Los textos entre comillas corresponden a botones '
        'o etiquetas tal como aparecen en pantalla.')


def s_acceso(c, solo_rol=None):
    d = c.doc
    d.add_heading('Acceso al sistema', 1)
    pasos(d, [
        'Abra el navegador e ingrese a la dirección del aplicativo (la URL interna que le '
        'indique el administrador).',
        'Escriba su Usuario y su Clave en el formulario de ingreso.',
        'Presione el botón "Ingresar".',
        'El sistema lo lleva automáticamente a su pantalla principal según su rol.',
    ])
    c.imagen('00_login.png',
             'Pantalla de inicio de sesión. El ojo junto al campo Clave permite ver lo que escribe.')
    if solo_rol:
        d.add_paragraph('Con su usuario de %s ingresará directamente a «%s».' % (solo_rol[0], solo_rol[1]))
    else:
        d.add_paragraph('La pantalla inicial que verá depende de su rol:')
        tabla(d,
              ['Rol', 'Pantalla inicial', 'Opciones del menú'],
              [
                  ['Supervisor', 'Solicitudes', 'Solicitudes (solo de su área)'],
                  ['Coordinador (jefe de transporte)', 'Coordinación',
                   'Coordinación, Dashboard, Reportes, Administración (catálogos limitados)'],
                  ['Conductor', 'Mis servicios', 'Mis servicios'],
                  ['Administrador', 'Coordinación',
                   'Solicitudes, Coordinación, Mis servicios, Dashboard, Reportes, Administración'],
              ],
              anchos=[4.2, 3.6, 8.7])
    d.add_paragraph(
        'Si el usuario o la clave son incorrectos aparece un aviso bajo el formulario. Verifique '
        'los datos o pida al administrador que revise su cuenta en Administración → Usuarios.')
    c.imagen('00b_login_error.png',
             'Aviso de error cuando el usuario o la clave no son válidos.')
    d.add_paragraph(
        'Para cerrar sesión presione "Cerrar sesión" (botón rojo al final del menú lateral) o el '
        'ícono de salida en el encabezado, junto a su nombre.')


def s_comunes(c, menu_items):
    d = c.doc
    d.add_heading('Elementos comunes de la pantalla', 1)
    d.add_paragraph(
        'Toda la aplicación comparte la misma estructura: a la izquierda el menú lateral verde '
        'con las opciones de su rol agrupadas por sección (Procesos, Reportes, Configuración); '
        'arriba el encabezado con el nombre del módulo, la fecha de hoy, su nombre con la '
        'etiqueta de su rol y los íconos de campana, llave y salir.')
    d.add_paragraph('Su menú muestra:')
    viñetas(d, menu_items)
    tabla(d,
          ['Elemento', 'Dónde está', 'Para qué sirve'],
          [
              ['Menú lateral', 'Columna izquierda',
               'Navegar entre los módulos de su rol. El botón ☰ lo pliega para ganar espacio.'],
              ['Campana', 'Encabezado, derecha',
               'Abre el panel de notificaciones. Un punto rojo indica avisos sin leer; toque una '
               'notificación para marcarla como leída.'],
              ['Badge amarillo', 'Encabezado, junto a la campana',
               'Aparece solo si hay operaciones guardadas sin internet pendientes de sincronizar.'],
              ['Llave', 'Encabezado, junto a su nombre',
               'Abre el panel "Cambiar clave": escriba la clave actual, la nueva (mínimo 6 '
               'caracteres) y su confirmación, luego "Guardar".'],
              ['Salir', 'Encabezado / menú lateral',
               'Cierra la sesión y vuelve a la pantalla de ingreso.'],
          ],
          anchos=[3.2, 4.3, 9.0])
    c.imagen('04_notificaciones.png',
             'Panel de notificaciones al presionar la campana. Las no leídas llegan también como '
             'aviso del navegador si dio el permiso.')
    c.imagen('42_cambiar_clave.png',
             'Panel "Cambiar clave" que se abre con el ícono de la llave del encabezado.')
    d.add_paragraph('El sistema le avisa por notificación cuando ocurren estos eventos:')
    viñetas(d, [
        'Le asignan una unidad a su solicitud (al supervisor que la creó y al conductor de la placa).',
        'Su solicitud se une o se acopla a un traslado, o es separada de uno.',
        'Su solicitud es desasignada de su unidad o cambia de estado (en ruta, realizada).',
        'El conductor agrega pasajeros de emergencia a un servicio.',
        'El coordinador publica un comunicado de "horas sin servicio".',
    ])

    d.add_heading('Aviso "Sin servicio" (comunicados)', 2)
    d.add_paragraph(
        'El coordinador puede publicar comunicados de suspensión del servicio por rango de '
        'horas (por ejemplo, cuando las unidades están en capacitación o mantenimiento). '
        'Si al ingresar hay un comunicado vigente, verá un aviso rojo bloqueante con el título, '
        'el horario afectado, las unidades ocupadas y el mensaje. Solo puede cerrarlo con el '
        'botón "Entendido". Mientras el rango esté vigente no se pueden crear solicitudes: '
        'la única excepción es una emergencia médica.')
    c.imagen('39_adm_comunicados.png',
             'Pestaña Comunicados en Administración, desde donde el coordinador o el '
             'administrador publican los avisos de suspensión del servicio.')

    d.add_heading('Tablas, paginador y botones con íconos', 2)
    viñetas(d, [
        '"Mostrando 1 a 10 de N registros": indica qué filas está viendo y el total.',
        'Flechas « ‹ › »: cambiar de página (primera, anterior, siguiente, última).',
        'Selector de filas: elija ver 5, 10, 20 o 50 registros por página.',
        'Los botones de acción de las tablas son solo íconos: en computadora pase el mouse para '
        'ver el nombre de la acción (tooltip); en celular el nombre aparece debajo del ícono.',
    ])
    tabla(d,
          ['Ícono', 'Acción', 'Dónde aparece'],
          [
              ['▶ (azul)', 'Iniciar el servicio', 'Mis servicios'],
              ['✓ (verde)', 'Finalizar el servicio', 'Mis servicios'],
              ['Persona+ (rojo)', 'Agregar pasajeros de emergencia', 'Mis servicios'],
              ['Bandera (celeste)', 'Abrir el checklist de paradas', 'Mis servicios'],
              ['Lápiz (azul)', 'Editar el registro', 'Solicitudes, Coordinación, Administración'],
              ['X (roja)', 'Anular o desactivar', 'Solicitudes, Coordinación, Administración'],
              ['⛓ (roja)', 'Separar del traslado', 'Coordinación'],
              ['⚡ (verde)', 'Asignación automática de unidad', 'Coordinación'],
              ['⧉ (celeste)', 'Asignación múltiple de unidades', 'Coordinación'],
              ['↩ (ámbar)', 'Desasignar la(s) unidad(es)', 'Coordinación'],
          ],
          anchos=[3.2, 6.5, 6.0])
    tabla(d,
          ['Color / etiqueta', 'Significado'],
          [
              ['Fila ámbar', 'Solicitud PENDIENTE: requiere atención del coordinador.'],
              ['Fila verde', 'Solicitud agrupada dentro de un traslado (T-N).'],
              ['Etiqueta ALTA (ámbar)', 'Prioridad alta: se atiende antes que las normales.'],
              ['Etiqueta EMERGENCIA (roja)', 'Prioridad máxima de atención.'],
              ['Etiqueta MÉDICA (roja)', 'El traslado responde a un caso médico.'],
              ['Etiqueta MULTIPLE', 'La solicitud está repartida entre varias unidades.'],
              ['Parada ámbar / azul / verde / gris', 'PENDIENTE / EN_PARADA / REALIZADA / OMITIDA.'],
          ],
          anchos=[5.5, 11.0])

    d.add_heading('Trabajar sin internet (modo offline)', 2)
    d.add_paragraph(
        'Si pierde la conexión, la aplicación sigue funcionando con la última información '
        'cargada: puede consultar listados que ya haya visto y ejecutar acciones (crear, '
        'asignar, marcar paradas, cambiar estado, etc.). Estas acciones se guardan en el equipo '
        'y aparecen contadas en el badge amarillo del encabezado. Al recuperar internet se '
        'envían solas, en el mismo orden, y el badge desaparece cuando todo quedó sincronizado. '
        'Si una acción ya no es válida al sincronizar (por ejemplo, el cupo se llenó mientras '
        'tanto), el sistema la descarta de la cola.')


def s_solicitudes(c):
    d = c.doc
    d.add_heading('Módulo Solicitudes (Supervisor)', 1)
    d.add_paragraph(
        'Es la pantalla principal del supervisor. Muestra únicamente las solicitudes de su área '
        '(el sistema la toma automáticamente de su usuario) y permite crear, editar, anular, '
        'acoplar a un traslado y hacer seguimiento en vivo.')
    c.imagen('01_sup_solicitudes.png',
             'Listado de solicitudes del supervisor: solo su área, con columnas PRIORIDAD, '
             'PLACA y CUPO de la unidad asignada.')

    d.add_heading('Entender el listado', 2)
    tabla(d,
          ['Columna', 'Qué muestra'],
          [
              ['PRIORIDAD', 'NORMAL, ALTA o EMERGENCIA, más la etiqueta MÉDICA si aplica.'],
              ['NOMBRE', 'Quién registró la solicitud.'],
              ['FECHA / HORA', 'Fecha y hora programada del viaje.'],
              ['ÁREA', 'Área del solicitante (siempre la suya).'],
              ['P. PARTIDA / P. LLEGADA', 'Punto de origen y de destino (del catálogo de Puntos).'],
              ['CANTIDAD', 'Número de personas a trasladar.'],
              ['MOTIVO', 'Motivo del viaje (capacitación, actividad, salud, etc.).'],
              ['PLACA', 'Unidad asignada, o MULTIPLE si se repartió entre varias (junto a ella se ven las placas con su cantidad).'],
              ['CUPO', 'Ocupadas/capacidad de la unidad. Verde = hay espacio, ámbar = casi llena, rojo = llena.'],
              ['ESTADO', 'PENDIENTE, ASIGNADO, EN_RUTA, REALIZADO o ANULADO.'],
              ['ACCIONES', 'Lápiz = editar (solo pendientes). X roja = anular (pide confirmación).'],
          ],
          anchos=[4.0, 12.5])
    d.add_paragraph(
        'Sobre la tabla hay dos campos de fecha "Desde" / "Hasta": al elegir un rango la lista '
        'se limita a esas fechas y aparece el botón "Ver todas" para quitar el filtro. Un clic '
        'sobre cualquier fila abre el detalle de la solicitud.')

    d.add_heading('Crear una solicitud', 2)
    pasos(d, [
        'Presione el botón "Crear" en la parte superior derecha.',
        'Su nombre y su área se cargan automáticamente (no se pueden modificar).',
        'La FECHA y la HORA PROGRAMADA se toman del sistema al momento de crear: la solicitud '
        'es para "ahora". Prográmela con la anticipación que requiera el traslado.',
        'Complete los campos obligatorios (marcados con *):',
    ])
    tabla(d,
          ['Campo', 'Obligatorio', 'Descripción'],
          [
              ['P. PARTIDA', 'Sí', 'Punto donde sube el personal (lista del catálogo de Puntos).'],
              ['P. LLEGADA', 'Sí', 'Punto de destino.'],
              ['CANTIDAD', 'Sí', 'Personas a trasladar (mínimo 1; flechas para subir/bajar).'],
              ['MOTIVO', 'Sí', 'Motivo del viaje (lista del catálogo de Motivos).'],
              ['OBSERVACIÓN', 'No', 'Comentario libre para el coordinador o el conductor.'],
              ['PRIORIDAD', 'Sí', 'Normal (por defecto), Alta o Emergencia. Define el orden de atención.'],
              ['Es emergencia médica', 'No', 'Casilla independiente de la prioridad: indica que el traslado es por un caso médico.'],
          ],
          anchos=[4.5, 2.5, 9.5])
    c.imagen('02_sup_crear.png',
             'Formulario de nueva solicitud: nombre y área fijos, fecha y hora automáticas.')
    c.imagen('02b_sup_crear_prioridad.png',
             'Parte inferior del formulario: OBSERVACIÓN, combo PRIORIDAD y la casilla '
             '"Es emergencia médica".')
    d.add_paragraph(
        'Al marcar "Es emergencia médica" el formulario se pre-llena solo: prioridad EMERGENCIA '
        '(bloqueada), cantidad 1, motivo "Salud" y destino "Tópico" si existen en los catálogos. '
        'Puede ajustar cantidad, motivo y destino si el caso lo requiere.')
    c.imagen('02c_sup_crear_emergencia.png',
             'Con "Es emergencia médica" marcada: prioridad fija en EMERGENCIA y sugerencias '
             'automáticas de motivo y destino.')
    pasos(d, [
        'Presione "Save". El botón "Cancel" cierra sin guardar.',
        'La solicitud aparece en el listado con estado PENDIENTE; el coordinador ya puede verla '
        'en Coordinación y usted recibirá una notificación cuando le asignen unidad.',
    ])
    d.add_paragraph(
        'Importante: si hay un comunicado "Sin servicio" vigente, el botón "Crear" abre el '
        'aviso rojo en lugar del formulario. Durante el bloqueo solo se admiten emergencias '
        'médicas, desde el botón "Es emergencia médica" del propio aviso.')

    d.add_heading('Ver el detalle y el seguimiento en vivo', 2)
    d.add_paragraph(
        'Haga clic sobre la fila de la solicitud. Se abre una ventana con todos sus datos: '
        'prioridad, solicitante, fecha y hora, área, puntos, cantidad, motivo, observación, '
        'placa y estado. Si la solicitud ya está dentro de un traslado, la ventana muestra '
        'además:')
    viñetas(d, [
        'La ruta del traslado y la última señal GPS de la unidad (se refresca cada 30 segundos '
        'mientras la ventana está abierta).',
        'La tabla de paradas con lo planeado (↑ suben / ↓ bajan), lo real que registró el '
        'conductor, el estado de cada parada y la hora.',
        'Las demás solicitudes que comparten el viaje.',
    ])
    c.imagen('03_sup_detalle.png',
             'Detalle de una solicitud con seguimiento en vivo: paradas con plan vs. real, '
             'estado y hora de cada una, y la señal GPS de la unidad.')

    d.add_heading('Acoplarse a una movilidad existente', 2)
    d.add_paragraph(
        'Si ya existe un traslado programado con cupo y una ruta compatible con su origen y '
        'destino, puede unir su solicitud a ese viaje sin esperar una nueva asignación. Solo '
        'puede acoplar solicitudes de su propia área.')
    pasos(d, [
        'Cree su solicitud y ábrala haciendo clic en su fila.',
        'Al final del detalle aparece la sección "Acoplar a traslado con cupo" con los '
        'traslados compatibles.',
        'Cada botón muestra el traslado (T-N), la placa y su ruta. Presione el que le sirva.',
        'El sistema verifica que haya cupo real para esa fecha; si lo hay, su solicitud queda '
        'ASIGNADA a ese traslado conservando su registro individual.',
    ])

    d.add_heading('Editar o anular una solicitud', 2)
    viñetas(d, [
        'Editar (lápiz azul): abre el formulario con los datos cargados para corregirlos. '
        'Disponible solo mientras la solicitud está PENDIENTE; una vez asignada, solo el '
        'coordinador puede modificarla o desasignarla.',
        'Anular (X roja): el sistema pide confirmación antes de cancelar. La acción no se puede '
        'deshacer y queda registrada en la auditoría.',
    ])

    d.add_heading('Estados de una solicitud', 2)
    tabla(d,
          ['Estado', 'Qué significa', 'Quién lo produce'],
          [
              ['PENDIENTE', 'Recién creada, aún sin unidad asignada.', 'Supervisor'],
              ['ASIGNADO', 'Ya tiene placa (sola, dentro de un traslado o repartida en varias unidades).', 'Coordinador / acople'],
              ['EN_RUTA', 'El conductor inició el recorrido.', 'Conductor'],
              ['REALIZADO', 'El conductor finalizó el viaje; se guardan las horas.', 'Conductor'],
              ['ANULADO', 'Solicitud cancelada.', 'Supervisor / coordinador'],
          ],
          anchos=[3.0, 9.0, 4.5])


def s_coordinacion(c):
    d = c.doc
    d.add_heading('Módulo Coordinación (Coordinador)', 1)
    d.add_paragraph(
        'Es el centro de operaciones del día. La pantalla se divide en dos paneles: a la '
        'izquierda las solicitudes con sus acciones, y a la derecha los traslados activos con '
        'su ocupación y seguimiento.')
    c.imagen('10_coo_coordinacion.png',
             'Coordinación del día: selector de fecha junto al título, badges de cupo por '
             'unidad, pestañas Activas/Realizadas/Anuladas y panel de Traslados activos.')

    d.add_heading('Filtro de fecha, pestañas y aviso de pendientes', 2)
    viñetas(d, [
        'Junto al título "Solicitudes" hay un selector de fecha (por defecto hoy). Todo lo que '
        'se ve corresponde a esa fecha: solicitudes, traslados y badges de cupo.',
        'El botón "Todas" muestra las solicitudes de todas las fechas y cambia a "Hoy" para '
        'volver. En modo "Todas" los badges de cupo se ocultan porque dependen de una fecha.',
        'Si hay solicitudes PENDIENTE en otras fechas aparece un aviso ámbar "Pendientes: '
        'dd/MM (N)". Presiónelo para saltar directo a la fecha pendiente más antigua.',
        'Las pestañas separan la vista: "Activas" (pendientes, asignadas y en ruta), '
        '"Realizadas" y "Anuladas", cada una con su contador.',
        'Las filas pendientes se pintan en ámbar para distinguirlas de un vistazo.',
    ])
    c.imagen('11_coo_todas.png',
             'Modo "Todas": historial completo de solicitudes sin filtro de fecha.')

    d.add_heading('Disponibilidad de las unidades (badges de cupo)', 2)
    d.add_paragraph(
        'Sobre la tabla hay un badge por cada combi con la forma "PLACA · libres/capacidad", '
        'por ejemplo "TP-001 · 15/15 libres". El cupo corresponde a la fecha seleccionada e '
        'incluye lo ocupado por asignaciones simples y por porciones de asignaciones múltiples:')
    viñetas(d, [
        'Verde: unidad libre o con bastante cupo.',
        'Ámbar: parcialmente ocupada (o no disponible).',
        'Rojo: llena — no acepta más pasajeros.',
    ])
    c.imagen('10b_coo_select_placa.png',
             'Selector de placa de una solicitud pendiente: cada opción muestra su cupo libre '
             'de la fecha; las unidades sin cupo quedan bloqueadas.')

    d.add_heading('Asignar unidad a una solicitud', 2)
    tabla(d,
          ['Botón', 'Acción'],
          [
              ['Selector "Placa" + ✓ azul', 'Elija la unidad en la lista (muestra el cupo de cada una) y presione el check para asignarla.'],
              ['⚡ (rayo verde)', 'Asignación automática: toma la primera unidad con cupo suficiente para esa fecha.'],
              ['⧉ (ventanas celeste)', 'Asignación múltiple: reparte la cantidad entre varias unidades cuando una sola no alcanza.'],
              ['Lápiz (azul)', 'Editar los datos de la solicitud (solo pendientes).'],
              ['X (roja)', 'Anular la solicitud (pide confirmación).'],
              ['↩ (ámbar)', 'Desasignar: quita la(s) unidad(es) de una solicitud ASIGNADA y la devuelve a PENDIENTE.'],
              ['⛓ (roja)', 'Separar del traslado (solo en solicitudes ya unidas y aún no iniciadas).'],
          ],
          anchos=[5.0, 11.5])
    pasos(d, [
        'En la columna ACCIONES de la solicitud elija la placa en el selector.',
        'Presione el botón ✓. El sistema valida el cupo real de la unidad para esa fecha.',
        'La solicitud pasa a ASIGNADO; el supervisor y el conductor de esa placa reciben una '
        'notificación.',
        'Para corregir una asignación use ↩ (Desasignar) y vuelva a asignar.',
    ])

    d.add_heading('Asignación múltiple (repartir entre unidades)', 2)
    d.add_paragraph(
        'Para solicitudes que superan la capacidad de una sola combi, presione ⧉. Se abre el '
        'panel "Distribuir N personas entre varias unidades":')
    pasos(d, [
        'En cada fila elija la unidad y escriba cuántas personas llevará.',
        'Con "Agregar unidad" sume tantas filas como necesite; el tachito quita una fila.',
        'El contador muestra "Asignado X / N personas": el total repartido debe coincidir '
        'exactamente con lo solicitado y cada unidad valida su cupo libre real.',
        'Presione "Guardar distribución".',
    ])
    c.imagen('15b_coo_multiple_2unid.png',
             'Diálogo de asignación múltiple: 20 personas repartidas entre TP-002 (15) y '
             'TP-003 (5).')
    d.add_paragraph(
        'En la tabla la solicitud muestra la etiqueta MULTIPLE con las placas y sus porciones '
        'en la columna TRASLADO; en el panel derecho aparece un traslado por cada unidad. Cada '
        'conductor ve solo la porción que le corresponde.')
    c.imagen('15c_coo_multiple_asignada.png',
             'Solicitud repartida: etiqueta MULTIPLE, porciones T-33 (TP-002) y T-34 (TP-003), '
             'y un traslado por unidad en el panel derecho.')

    d.add_heading('Unir solicitudes en un traslado', 2)
    pasos(d, [
        'Marque con los checkboxes de la primera columna dos o más solicitudes pendientes de '
        'la misma fecha (solo las pendientes se pueden seleccionar).',
        'Arriba de la tabla aparecen el total de personas, el selector "Placa del traslado" y '
        'el botón "Unir (N)".',
        'Elija la placa y presione "Unir". El sistema valida el cupo: si la suma excede lo '
        'disponible, el botón se bloquea.',
        'Se crea el traslado T-N que agrupa las solicitudes: cada una conserva su número, '
        'solicitante y cantidad para los reportes.',
    ])
    c.imagen('10c_coo_unir.png',
             'Dos solicitudes seleccionadas: aparecen el total de personas, el selector de '
             'placa y el botón "Unir (2)".')
    c.imagen('10e_coo_traslado_creado.png',
             'Traslado T-32 creado: las solicitudes quedan ASIGNADO (verde) y la tarjeta del '
             'traslado aparece en el panel derecho con su ocupación.')
    d.add_paragraph(
        'Separar: en la fila de una solicitud dentro de un traslado aparece el botón ⛓. Tras '
        'confirmar, la solicitud vuelve a PENDIENTE sin placa y el traslado recalcula su '
        'ocupación; si queda vacío se anula solo. Un clic sobre la etiqueta T-N de la tabla '
        'lleva la página del panel derecho a ese traslado y lo resalta.')

    d.add_heading('Definir las paradas y la ruta', 2)
    pasos(d, [
        'En la tarjeta del traslado (panel derecho) presione "Editar paradas".',
        'Con "Agregar parada" añada filas: en cada una elija el punto y escriba cuántas personas '
        'suben (↑) y cuántas bajan (↓) ahí.',
        'Use las flechas ↑ ↓ para ordenar la secuencia de recogida y el tachito para quitar una '
        'parada.',
        'El encabezado muestra el balance "Total: ↑ N suben · ↓ N bajan": lo que sube en toda '
        'la ruta debe bajar completo. Al editar una bajada intermedia, la última parada absorbe '
        'la diferencia automáticamente.',
        'El botón "Ruta sugerida" arma la secuencia solo: fusiona las partidas y llegadas de '
        'las solicitudes del traslado ordenadas por hora, con el detalle de qué solicitud sube '
        'y baja en cada punto.',
        'Presione "Guardar ruta". El conductor verá las paradas en ese orden en su checklist.',
    ])
    c.imagen('12_coo_editor_paradas.png',
             'Editor de ruta del traslado: guía de uso, balance de subidas/bajadas, botón '
             '"Ruta sugerida" y una fila por parada con su orden.')
    d.add_paragraph(
        'Una vez que el conductor marcó la primera parada (ruta en ejecución) el editor pasa a '
        'solo lectura ("Ver ruta"): los ajustes en caliente los hace el propio conductor con '
        '"Parada imprevista" en su aplicación.')

    d.add_heading('Seguimiento en vivo del traslado', 2)
    d.add_paragraph(
        'En la tarjeta del traslado, "Progreso" abre la línea de tiempo de las paradas y se '
        'actualiza sola cada 30 segundos. La tabla muestra por cada parada: lo planeado '
        '(↑ sube / ↓ baja), lo real registrado por el conductor, cuántos quedan a bordo, el '
        'estado (PENDIENTE, EN_PARADA, REALIZADA, OMITIDA), las horas de llegada y salida, y el '
        'detalle de qué solicitud sube o baja en ese punto. El encabezado resume "x/y paradas · '
        'N a bordo" y la hora de la última señal GPS que reportó la unidad.')
    c.imagen('13_coo_progreso.png',
             'Progreso del traslado T-32: línea de tiempo de paradas con plan vs. real, '
             'pasajeros a bordo, estados y horas.')
    d.add_paragraph(
        'Los traslados del día que ya terminaron pasan a la sección plegable "Finalizados", '
        'debajo de los activos: desde ahí puede consultar su progreso y ruta en solo lectura.')
    c.imagen('19_coo_finalizados.png',
             'Sección "Finalizados" expandida: traslados cerrados del período consultado, '
             'solo para consulta.')

    d.add_heading('Editar y anular solicitudes', 2)
    d.add_paragraph(
        'El lápiz abre un panel con todos los campos de la solicitud pendiente: fecha, hora, '
        'puntos, cantidad, motivo, observación, prioridad y la marca de emergencia médica. '
        'Todo cambio queda registrado en la auditoría.')
    c.imagen('16_coo_editar_solicitud.png',
             'Edición de una solicitud pendiente desde Coordinación.')


def s_conductor(c):
    d = c.doc
    d.add_heading('Módulo Mis servicios (Conductor)', 1)
    d.add_paragraph(
        'El conductor solo ve los servicios asignados a la placa de su unidad, ordenados con '
        'las emergencias primero. Arriba hay un selector de fecha (por defecto hoy) con los '
        'botones "Todas" / "Hoy" y "Actualizar", y la franja "Ruta del día" con la secuencia de '
        'puntos del día seleccionado.')
    c.imagen('20_con_servicios.png',
             'Mis servicios: banner "Ruta del día", servicio ASIGNADO con su cupo (15/15), '
             'indicador de paradas 0/2 y los botones de acción.')
    tabla(d,
          ['Botón (columna ACCIONES)', 'Qué hace'],
          [
              ['▶ (azul)', 'Iniciar el servicio: aparece en estado ASIGNADO. Pasa a EN_RUTA y guarda la hora de inicio. Desde ese momento la unidad reporta su GPS cada 30 segundos.'],
              ['✓ (verde)', 'Finalizar: aparece en EN_RUTA. Pasa a REALIZADO y guarda la hora de fin. Solo se habilita cuando todas las paradas ya están registradas u omitidas.'],
              ['Persona+ (rojo)', 'Agregar pasajeros de emergencia si hay cupo libre.'],
              ['Bandera (celeste)', 'Abrir el checklist de paradas del servicio.'],
          ],
          anchos=[5.0, 11.5])
    pasos(d, [
        'Al llegar a su turno, abra "Mis servicios" y revise los viajes del día: origen, '
        'destino, cantidad de personas, motivo y si es emergencia médica.',
        'Si un servicio es parte de una solicitud repartida en varias unidades, verá solo su '
        'porción con la cantidad que le corresponde a su combi.',
        'Presione la bandera para conocer las paradas ordenadas antes de arrancar.',
        'Al iniciar el recorrido presione ▶. El estado cambia a EN_RUTA y su unidad empieza a '
        'reportar GPS.',
        'Recorra las paradas en el orden indicado registrando cada una (ver siguiente sección).',
        'Al terminar presione ✓ (Finalizar). En solicitudes repartidas, cada conductor finaliza '
        'su porción; la solicitud se marca REALIZADO cuando todas las porciones terminaron.',
        'Sin señal puede seguir operando: los cambios se guardan en el equipo y se envían al '
        'recuperar internet (con la hora real en que ocurrieron).',
    ])

    d.add_heading('Checklist de paradas', 2)
    d.add_paragraph(
        'El botón de la bandera abre el panel con la ruta programada: cada parada muestra su '
        'tipo (PARTIDA, PARADA o LLEGADA), lo planeado (↑ suben / ↓ bajan), lo real, cuántos '
        'quedan a bordo, el estado y la hora. Solo se puede actuar sobre la siguiente parada '
        'pendiente, en orden:')
    pasos(d, [
        'Con el servicio EN_RUTA, al llegar al punto presione el botón de llegada de la parada '
        '(azul): queda EN_PARADA con la hora y el GPS del momento.',
        'Presione "Registrar" (verde) y confirme cuántos subieron y cuántos bajaron realmente; '
        'viene prellenado con lo planeado para corregir solo lo que cambió.',
        'Si la parada no tuvo movimiento (nadie se presentó, punto equivocado), use "Omitir" y '
        'escriba el motivo: queda registrado para el coordinador y la auditoría.',
        'Si debe recoger a alguien en un punto no programado, use "+ Parada imprevista": elija '
        'el punto y cuántos suben; se inserta como la siguiente parada del recorrido.',
    ])
    c.imagen('21_con_paradas.png',
             'Checklist de paradas del servicio: tipo de punto, plan vs. real, pasajeros a '
             'bordo y estado de cada parada.')
    d.add_paragraph(
        'La app no deja finalizar el servicio si quedan paradas sin procesar: el botón ✓ avisa '
        'cuáles faltan para que las registre o las omita primero.')

    d.add_heading('Pasajeros de emergencia', 2)
    pasos(d, [
        'En un servicio ASIGNADO o EN_RUTA presione el botón rojo de persona+.',
        'El panel muestra el cupo libre de la unidad (ocupados/capacidad del día).',
        'Escriba cuántos pasajeros adicionales suben y presione "Agregar". Si el servicio tiene '
        'paradas, el extra queda amarrado a la parada actual.',
        'Si la cantidad supera el cupo disponible el sistema lo rechaza. Al guardar, esa '
        'porción queda marcada como emergencia y el supervisor recibe una notificación.',
    ])
    c.imagen('22_con_pasajeros.png',
             'Panel para agregar pasajeros de emergencia con el cupo libre de la unidad.')


def s_dashboard(c):
    d = c.doc
    d.add_heading('Módulo Dashboard (Coordinador / Administrador)', 1)
    d.add_paragraph(
        'Tablero visual de la operación. Elija el rango "Desde" / "Hasta" (por defecto el mes '
        'en curso) y presione "Generar".')
    c.imagen('30_adm_dashboard.png',
             'Dashboard: rango de fechas, tabla de KPIs (viajes, personal planificado vs. real '
             'según paradas, motivos principales) y gráficos.')
    d.add_paragraph('La tabla de KPIs resume:')
    viñetas(d, [
        'Cantidad de viajes en el rango.',
        'Personal transportado planificado (lo que se solicitó).',
        'Personal transportado real (conteo efectivo registrado en las paradas por el '
        'conductor) y su porcentaje de cobertura.',
        'Los motivos de movimiento más frecuentes con su participación.',
    ])
    d.add_paragraph('Debajo se muestran ocho gráficos:')
    viñetas(d, [
        'Movimiento de personal por cada área (barras).',
        'Motivo de movimiento (barras).',
        'Total de viajes realizados vs. no realizados (torta).',
        'Horas con mayor frecuencia de movimiento.',
        'Personal trasladado por unidad.',
        'Áreas que más personal trasladaron (top 8).',
        'Personal transportado por semana.',
        'Tendencia por mes (línea).',
    ])
    c.imagen('30b_adm_dashboard2.png', 'Gráficos del dashboard (parte media).')
    c.imagen('30c_adm_dashboard3.png', 'Gráficos del dashboard (parte final).')


def s_reportes(c):
    d = c.doc
    d.add_heading('Módulo Reportes (Coordinador / Administrador)', 1)
    d.add_paragraph('Genera el histórico de movilidades y los indicadores de operación.')
    pasos(d, [
        'Elija el rango "Desde" / "Hasta" y presione "Generar" (por defecto el mes en curso).',
        'Opcionalmente refine con los filtros Supervisor, Placa y Estado.',
        'Revise las tarjetas de indicadores y las tablas de eficiencia.',
        'Use "Exportar" para descargar el detalle: CSV (compatible con Excel) o Excel con el '
        'formato de reporte de movilidad (columnas MES, SEMANA, DIA, …, INTERVALO). Con '
        '"Imprimir / PDF" obtiene la versión imprimible.',
    ])
    c.imagen('31_adm_reportes.png',
             'Reportes: filtros de fecha y supervisor/placa/estado, totales y tarjetas de '
             'indicadores.')
    tabla(d,
          ['Indicador', 'Qué mide'],
          [
              ['Solicitudes / Personas', 'Total de movilidades y de personas trasladadas en el rango.'],
              ['Realizados / Emergencias', 'Viajes completados y cuántos fueron de emergencia.'],
              ['Duración promedio', 'Tiempo medio entre Iniciar y Finalizar.'],
              ['Espera promedio', 'Tiempo medio entre la creación de la solicitud y su atención.'],
              ['Puntualidad ±10 min', 'Servicios realizados dentro de 10 minutos de la hora programada.'],
              ['Eficiencia por unidad', 'Viajes, personas y ocupación por combi (incluye las porciones de asignaciones múltiples).'],
              ['Servicios por conductor', 'Viajes y personas atendidas por cada conductor.'],
              ['Movimiento por área', 'Solicitudes y personas agrupadas por el área del maestro.'],
          ],
          anchos=[5.0, 11.5])
    c.imagen('31b_adm_reporte_tabla.png',
             'Detalle del reporte: una fila por solicitud con prioridad, placa (o MULTIPLE con '
             'sus placas) y estado.')
    d.add_paragraph(
        'La columna de prioridad del detalle muestra NORMAL / ALTA / EMERGENCIA y la etiqueta '
        'MÉDICA cuando corresponde; ambas se exportan.')

    d.add_heading('Auditoría', 2)
    d.add_paragraph(
        'El botón "Cargar auditoría" muestra el historial de todas las acciones del sistema: '
        'inicios de sesión, creación y edición de solicitudes, asignaciones, uniones, '
        'separaciones, desasignaciones, cambios de estado, registro de paradas y pasajeros '
        'agregados, con fecha, entidad afectada y usuario. El mismo historial está disponible '
        'en la pestaña Auditoría de Administración (solo administrador).')
    d.add_paragraph(
        'La columna Detalle guarda el JSON exacto de la operación (qué se envió al sistema), '
        'útil para rastrear por ejemplo qué placas y cantidades se repartieron en una '
        'asignación múltiple.')


def s_administracion(c, es_coordinador=False):
    d = c.doc
    titulo = 'Módulo Administración'
    d.add_heading(titulo + (' (Coordinador — vista limitada)' if es_coordinador else ' (Administrador)'), 1)
    if es_coordinador:
        d.add_paragraph(
            'El coordinador ve un subconjunto de pestañas: Usuarios, Conductores, Unidades, '
            'Tipos vehículo, Puntos, Áreas y Comunicados. Puede crear usuarios únicamente con '
            'rol Supervisor y gestionar solo registros de supervisores y conductores (no otros '
            'coordinadores ni administradores). En Áreas no ve las columnas contables '
            '(CeCo/cuenta/proyecto), que son exclusivas del administrador.')
    else:
        d.add_paragraph(
            'Gestiona los catálogos, las unidades, los comunicados y la auditoría del sistema '
            'en nueve pestañas.')
    d.add_paragraph(
        'Todas las pestañas de catálogo funcionan igual: el formulario de arriba sirve para '
        'crear (botón "Crear") y para editar — al presionar el lápiz de una fila sus datos se '
        'cargan en el formulario y el botón cambia a "Actualizar" junto a "Cancelar". El botón '
        'de prohibido desactiva el registro sin borrarlo (en usuarios también existe el '
        'tachito, que es un borrado lógico: la cuenta ya no puede iniciar sesión pero su '
        'historial se conserva; la casilla "Mostrar eliminados" los vuelve visibles). En '
        'usuarios y conductores, si deja la clave vacía al editar se conserva la actual.')
    c.imagen('41_adm_editar.png',
             'Edición en Administración: el lápiz carga el registro en el formulario y el '
             'botón pasa a "Actualizar".')
    tabla(d,
          ['Pestaña', 'Campos / contenido', 'Para qué sirve'],
          [
              ['Usuarios', 'Usuario, nombre, DNI, rol, área (del maestro) y clave.',
               'Cuentas de supervisores, coordinadores y administradores.'],
              ['Conductores', 'Usuario, nombre, DNI, placa de la unidad asignada, área y clave.',
               'Cuentas de conductor; la placa define qué servicios verá en "Mis servicios".'],
              ['Unidades', 'Placa, propiedad (Propia/Tercero), tipo de vehículo, capacidad, '
               'estado operativo (Operativa / Mantenimiento / Fuera de servicio), conductor '
               'asignado, marca, modelo y año.',
               'Las combis disponibles para asignar; el conductor asignado recibe los servicios '
               'de esa placa.'],
              ['Tipos vehículo', 'Nombre.',
               'Catálogo que alimenta el combo "Tipo de vehículo" de Unidades.'],
              ['Puntos', 'Nombre, latitud y longitud.',
               'Orígenes/destinos y paradas. Las coordenadas alimentan el mapa y la validación '
               'de rutas al acoplarse.'],
              ['Motivos', 'Nombre.',
               'Catálogo de motivos de viaje que elige el supervisor.'],
              ['Áreas', 'Nombre' + ('' if es_coordinador else ' (y para el administrador: CeCo, cuenta y proyecto).'),
               'Maestro de áreas: se asigna a usuarios y conductores; define qué solicitudes ve '
               'cada supervisor.'],
              ['Comunicados', 'Título, fecha, hora desde/hasta, actividad, unidades ocupadas '
               '(todas o una) y mensaje.',
               'Publica los rangos "sin servicio": bloquean la creación de solicitudes (salvo '
               'emergencias médicas) y muestran el aviso rojo a todos los usuarios.'],
              ['Auditoría', 'Fecha, entidad, acción, usuario y detalle (solo lectura).',
               'Historial completo de todas las acciones del sistema.'],
          ],
          anchos=[2.8, 6.9, 6.8])
    c.imagen('32_adm_usuarios.png', 'Pestaña Usuarios: formulario arriba y tabla con estado.')
    c.imagen('33_adm_conductores.png', 'Pestaña Conductores: usuario, DNI, placa asignada y área.')
    c.imagen('34_adm_unidades.png',
             'Pestaña Unidades: propiedad, tipo de vehículo, capacidad, conductor asignado y '
             'estado operativo.')
    c.imagen('36_adm_puntos.png', 'Pestaña Puntos: nombre con latitud y longitud.')
    c.imagen('38_adm_areas.png', 'Pestaña Áreas: catálogo maestro que se asigna a usuarios y conductores.')

    d.add_heading('Comunicados "sin servicio"', 2)
    pasos(d, [
        'Entre a la pestaña Comunicados y complete: título, fecha afectada, hora de inicio y '
        'fin del bloqueo, actividad (capacitación, mantenimiento…) y si ocupa todas las '
        'unidades o una en particular.',
        'Escriba el mensaje que verán los usuarios y presione "Crear".',
        'Mientras el rango esté vigente, todos los usuarios ven el aviso rojo al entrar y no '
        'se pueden crear solicitudes (solo emergencias médicas). En la tabla, la columna '
        '"Vigente" indica si está dentro del horario.',
        'Para levantar el bloqueo antes de tiempo, desactive el comunicado con el botón de '
        'prohibido.',
    ])
    c.imagen('39_adm_comunicados.png',
             'Pestaña Comunicados: formulario de publicación y listado con su vigencia.')

    if not es_coordinador:
        c.imagen('40_adm_auditoria.png',
                 'Pestaña Auditoría: historial de acciones con fecha, entidad, usuario y '
                 'detalle de la operación.')
        d.add_paragraph(
            'Importante sobre las áreas: asígnele un área a cada supervisor desde el combo (no '
            'es texto libre). Un supervisor sin área solo verá las solicitudes que él mismo '
            'registró. El área de la solicitud se toma automáticamente del supervisor que la crea.')


def s_glosario(c):
    d = c.doc
    d.add_heading('Glosario', 1)
    tabla(d,
          ['Término', 'Significado'],
          [
              ['Solicitud', 'Pedido de transporte creado por un supervisor para su área.'],
              ['Traslado (T-N)', 'Agrupación de una o más solicitudes bajo una misma placa, con su ruta y paradas ordenadas.'],
              ['Acoplarse', 'Unir una solicitud pendiente a un traslado con cupo y ruta compatible.'],
              ['Asignación múltiple', 'Reparto de una solicitud entre varias unidades; la etiqueta MULTIPLE lo indica.'],
              ['Porción', 'Parte de una solicitud múltiple que atiende una unidad concreta.'],
              ['Parada', 'Punto de la ruta donde sube o baja personal; el conductor la marca como EN_PARADA, REALIZADA u OMITIDA.'],
              ['Parada imprevista', 'Parada no programada que el conductor inserta en ruta (recojo extra, desvío).'],
              ['Cupo', 'Asientos libres de la unidad en una fecha (ocupadas/capacidad).'],
              ['Prioridad', 'Nivel de urgencia de la solicitud: NORMAL, ALTA o EMERGENCIA.'],
              ['Emergencia médica', 'Marca que indica que el traslado responde a un caso médico; es la única que pasa durante un comunicado "sin servicio".'],
              ['Comunicado', 'Aviso de suspensión del servicio por rango de horas que publica el coordinador.'],
              ['Área', 'Unidad organizativa del supervisor; define qué solicitudes ve y cómo se agrupan los reportes.'],
              ['Badge', 'Etiqueta de color que resume un estado (verde/ámbar/rojo).'],
              ['Tooltip', 'Etiqueta flotante con el nombre de un botón; aparece al pasar el mouse.'],
              ['Paginador', 'Controles al pie de cada tabla para cambiar de página y elegir cuántas filas ver.'],
              ['Auditoría', 'Registro de quién hizo cada cambio y cuándo, con el detalle de la operación.'],
              ['Modo offline', 'Trabajo sin internet: las acciones se guardan en el equipo y se sincronizan al volver la conexión.'],
          ],
          anchos=[3.5, 13.0])


# ===================== documento general =====================

def manual_general():
    c = Ctx()
    d = c.doc
    portada(d, 'Manual de Usuario', 'Aplicativo de Transporte de Personal — Manual General')
    indice(d)
    s_intro(c, 'Este manual cubre todas las funciones del sistema para los cuatro roles: '
               'Supervisor, Coordinador, Conductor y Administrador.')
    s_acceso(c)
    s_comunes(c, [
        'Supervisor: Solicitudes.',
        'Coordinador: Coordinación, Dashboard, Reportes, Administración (catálogos limitados).',
        'Conductor: Mis servicios.',
        'Administrador: Solicitudes, Coordinación, Mis servicios, Dashboard, Reportes y '
        'Administración completa.',
    ])
    s_solicitudes(c)
    s_coordinacion(c)
    s_conductor(c)
    s_dashboard(c)
    s_reportes(c)
    s_administracion(c, es_coordinador=False)

    d.add_heading('Guía rápida por rol', 1)
    tabla(d,
          ['Rol', 'Flujo típico'],
          [
              ['Supervisor',
               '1. Ingresar → Solicitudes (solo su área). 2. "Crear", elegir prioridad y marcar '
               'emergencia médica si aplica. 3. (Opcional) Abrir el detalle y acoplarse a un '
               'traslado con cupo. 4. Seguir el ESTADO y el seguimiento en vivo. 5. Editar o '
               'anular mientras esté pendiente.'],
              ['Coordinador',
               '1. Ingresar → Coordinación. 2. Elegir la fecha y atender el aviso de pendientes. '
               '3. Asignar placa (✓, ⚡ o ⧉) o unir solicitudes compatibles en un traslado. '
               '4. "Editar paradas" para ordenar el recorrido (o "Ruta sugerida"). '
               '5. Seguir el "Progreso" en vivo y el GPS. 6. Desasignar (↩) o separar (⛓) para '
               'corregir. 7. Publicar comunicados si el servicio se suspende. 8. Revisar '
               'Dashboard y Reportes al cierre.'],
              ['Conductor',
               '1. Ingresar → Mis servicios (fecha de hoy o "Todas"). 2. Revisar la "Ruta del '
               'día" y las paradas con la bandera. 3. ▶ al arrancar. 4. En cada parada: '
               'llegada → registrar suben/bajan (u omitir con motivo). 5. Persona+ solo para '
               'pasajeros de emergencia con cupo. 6. ✓ al terminar su porción del viaje.'],
              ['Administrador',
               '1. Mantener catálogos: usuarios con su área, conductores con su placa, unidades '
               'con su conductor, tipos de vehículo, puntos con coordenadas, motivos y áreas. '
               '2. Publicar comunicados. 3. Revisar Dashboard, Reportes y Auditoría. '
               '4. Puede operar cualquier pantalla de los demás roles.'],
          ],
          anchos=[3.5, 13.0])
    s_glosario(c)
    c.save('Manual_General_Transporte_Personal_V4.docx')


# ===================== manuales por rol =====================

def manual_supervisor():
    c = Ctx()
    d = c.doc
    portada(d, 'Manual del Supervisor', 'Aplicativo de Transporte de Personal — Rol Supervisor')
    indice(d)
    s_intro(c, 'Este manual cubre el rol Supervisor: crear y gestionar las solicitudes de '
               'movilidad de su área, acoplarse a traslados existentes y hacer seguimiento en '
               'vivo del viaje.')
    s_acceso(c, ('Supervisor', 'Solicitudes'))
    s_comunes(c, ['Solicitudes: su único módulo, con las movilidades de su área.'])
    s_solicitudes(c)
    d.add_heading('Resumen de su flujo de trabajo', 1)
    pasos(d, [
        'Ingrese con su usuario: caerá directamente en Solicitudes.',
        'Presione "Crear" y complete punto de partida, punto de llegada, cantidad y motivo. '
        'La fecha y hora se toman del sistema.',
        'Elija la prioridad (Normal/Alta/Emergencia) y marque "Es emergencia médica" si '
        'corresponde.',
        'Si ya hay un traslado compatible, abra el detalle de su solicitud y acóplese.',
        'Siga el estado: PENDIENTE → ASIGNADO → EN_RUTA → REALIZADO. Recibirá notificaciones '
        'en cada cambio.',
        'Mientras esté PENDIENTE puede editarla (lápiz) o anularla (X).',
        'Si el coordinador publicó un comunicado "sin servicio", no podrá crear solicitudes '
        'en ese horario salvo emergencias médicas.',
    ])
    s_glosario(c)
    c.save('Manual_Rol_Supervisor.docx')


def manual_coordinador():
    c = Ctx()
    d = c.doc
    portada(d, 'Manual del Coordinador', 'Aplicativo de Transporte de Personal — Rol Coordinador (jefe de transporte)')
    indice(d)
    s_intro(c, 'Este manual cubre el rol Coordinador (jefe de transporte): asignar unidades a '
               'las solicitudes, unir y separar traslados, definir paradas, hacer seguimiento '
               'en vivo, publicar comunicados y revisar dashboard, reportes y catálogos.')
    s_acceso(c, ('Coordinador', 'Coordinación'))
    s_comunes(c, [
        'Coordinación: centro de operaciones (solicitudes + traslados).',
        'Dashboard: gráficos de la operación.',
        'Reportes: histórico, indicadores y auditoría.',
        'Administración: catálogos limitados (usuarios supervisores, conductores, unidades, '
        'tipos de vehículo, puntos, áreas y comunicados).',
    ])
    s_coordinacion(c)
    s_dashboard(c)
    s_reportes(c)
    s_administracion(c, es_coordinador=True)
    d.add_heading('Resumen de su flujo de trabajo', 1)
    pasos(d, [
        'Ingrese → Coordinación, con la fecha de hoy cargada.',
        'Atienda el aviso ámbar de pendientes si aparece (salta a la fecha más antigua).',
        'Asigne unidad a cada solicitud pendiente: selector + ✓, ⚡ automático o ⧉ múltiple '
        'si no cabe en una combi.',
        'Una solicitudes compatibles en un traslado (checkboxes → placa → "Unir").',
        'Defina la ruta con "Editar paradas" o "Ruta sugerida" y guárdela.',
        'Durante el día siga "Progreso" en vivo y la última señal GPS de cada unidad.',
        'Corrija con ↩ (desasignar) o ⛓ (separar) cuando haga falta.',
        'Publique comunicados en Administración → Comunicados cuando el servicio se suspenda.',
        'Cierre el día revisando Dashboard y Reportes.',
    ])
    s_glosario(c)
    c.save('Manual_Rol_Coordinador.docx')


def manual_conductor():
    c = Ctx()
    d = c.doc
    portada(d, 'Manual del Conductor', 'Aplicativo de Transporte de Personal — Rol Conductor')
    indice(d)
    s_intro(c, 'Este manual cubre el rol Conductor: ver los servicios asignados a su unidad, '
               'iniciarlos, registrar cada parada del recorrido, agregar pasajeros de '
               'emergencia y finalizar el viaje.')
    s_acceso(c, ('Conductor', 'Mis servicios'))
    s_comunes(c, ['Mis servicios: su único módulo, con los viajes de su placa.'])
    s_conductor(c)
    d.add_heading('Resumen de su jornada', 1)
    pasos(d, [
        'Ingrese con su usuario de conductor: caerá directamente en Mis servicios con la fecha '
        'de hoy.',
        'Revise la "Ruta del día" y abra el checklist de paradas (bandera) para conocer el '
        'recorrido.',
        'Al arrancar, presione ▶ (Iniciar). Desde ahí su unidad reporta GPS cada 30 segundos.',
        'En cada parada: presione llegada al llegar, luego "Registrar" con los conteos reales '
        'de subidas y bajadas. Si la parada no tiene movimiento, "Omitir" con su motivo.',
        'Para un recojo no previsto use "+ Parada imprevista".',
        'Con persona+ agregue pasajeros de emergencia solo si hay cupo.',
        'Al terminar todas las paradas presione ✓ (Finalizar).',
        'Sin internet siga marcando: todo se sincroniza solo cuando vuelva la señal.',
    ])
    s_glosario(c)
    c.save('Manual_Rol_Conductor.docx')


def manual_administrador():
    c = Ctx()
    d = c.doc
    portada(d, 'Manual del Administrador', 'Aplicativo de Transporte de Personal — Rol Administrador')
    indice(d)
    s_intro(c, 'Este manual cubre el rol Administrador: ve y opera todos los módulos del '
               'sistema (Solicitudes, Coordinación, Mis servicios, Dashboard, Reportes y la '
               'Administración completa de catálogos, comunicados y auditoría).')
    s_acceso(c, ('Administrador', 'Coordinación'))
    s_comunes(c, [
        'Solicitudes, Coordinación y Mis servicios: los módulos operativos de los otros roles.',
        'Dashboard y Reportes: indicadores, gráficos, exportaciones y auditoría.',
        'Administración: las nueve pestañas de catálogos, comunicados y auditoría.',
    ])
    s_solicitudes(c)
    s_coordinacion(c)
    s_conductor(c)
    s_dashboard(c)
    s_reportes(c)
    s_administracion(c, es_coordinador=False)
    d.add_heading('Resumen de sus tareas', 1)
    viñetas(d, [
        'Mantener usuarios: cada supervisor con su área; cada conductor con su placa.',
        'Mantener la flota: unidades con su tipo, capacidad, estado operativo y conductor '
        'asignado; catálogo de tipos de vehículo.',
        'Mantener puntos con coordenadas (alimentan mapas y acoples), motivos y áreas.',
        'Publicar o desactivar comunicados de suspensión del servicio.',
        'Supervisar la operación desde Dashboard y Reportes, y auditar acciones en Auditoría.',
        'Operar cualquier módulo de los otros roles (solicita, coordina o conduce con su '
        'cuenta de administrador).',
    ])
    s_glosario(c)
    c.save('Manual_Rol_Administrador.docx')


if __name__ == '__main__':
    manual_general()
    manual_supervisor()
    manual_coordinador()
    manual_conductor()
    manual_administrador()

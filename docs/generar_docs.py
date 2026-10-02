# -*- coding: utf-8 -*-
"""Genera los documentos Word del sistema de Transporte de Personal."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

MORADO = RGBColor(0x7B, 0x1F, 0xA2)
GRIS = RGBColor(0x55, 0x55, 0x55)
OUT = os.path.dirname(os.path.abspath(__file__))


def estilo_base(doc):
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    for h, sz in (('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12)):
        s = doc.styles[h]
        s.font.name = 'Calibri'
        s.font.size = Pt(sz)
        s.font.color.rgb = MORADO
        s.font.bold = True


def portada(doc, titulo, subtitulo):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo)
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = MORADO
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitulo)
    r2.font.size = Pt(14)
    r2.font.color.rgb = GRIS
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run('\nAplicativo de Transporte de Personal Interno\nVersión 1.0 — Septiembre 2026').font.color.rgb = GRIS
    doc.add_page_break()


def tabla(doc, cabecera, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabecera):
        cell = t.rows[0].cells[i]
        cell.text = c
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
    for fila in filas:
        cells = t.add_row().cells
        for i, v in enumerate(fila):
            cells[i].text = str(v)
    if anchos:
        for i, w in enumerate(anchos):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def viñetas(doc, items):
    for it in items:
        doc.add_paragraph(it, style='List Bullet')


def pasos(doc, items):
    for it in items:
        doc.add_paragraph(it, style='List Number')


def imagen(doc, archivo, caption, ancho=Cm(15.5)):
    """Inserta una captura de docs/img centrada con pie de figura."""
    ruta = os.path.join(OUT, 'img', archivo)
    if not os.path.exists(ruta):
        print('  (falta imagen: ' + archivo + ')')
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=ancho)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = GRIS


# ============================================================
# DOCUMENTO 1: FLUJO DEL SISTEMA Y USO POR ROL
# ============================================================
d = Document()
estilo_base(d)
portada(d, 'Flujo del Sistema y Uso por Rol', 'Documento funcional — Transporte de Personal')

d.add_heading('1. Objetivo del sistema', 1)
d.add_paragraph(
    'El aplicativo de Transporte de Personal digitaliza la gestión de solicitudes de movilidad '
    'interna, reemplazando el manejo por WhatsApp y hojas de cálculo. Cubre el ciclo completo: '
    'creación de la solicitud, asignación de unidades (combis), definición de rutas con paradas '
    'intermedias, ejecución del viaje por el conductor y generación de reportes y auditoría.')
d.add_paragraph(
    'La flota actual consta de 3 combis internas con capacidad aproximada de 15 personas cada una. '
    'Las solicitudes de gran volumen (50-100+ personas) se siguen gestionando con buses externos; '
    'este aplicativo cubre la movilidad interna de grupos pequeños.')

d.add_heading('2. Roles del sistema', 1)
d.add_paragraph('El sistema maneja cuatro roles. El menú lateral y las pantallas disponibles cambian según el rol del usuario autenticado.')
tabla(d,
      ['Rol', 'Código', 'Qué puede hacer', 'Pantallas'],
      [
          ['Supervisor', 'SPTRANS',
           'Crear solicitudes de movilidad, ver el listado programado, acoplarse a traslados con cupo, editar o anular sus solicitudes.',
           'Solicitudes'],
          ['Jefe de Transporte (Coordinador)', 'COTRANS',
           'Ver todas las solicitudes, asignar placas (manual, automática o múltiple), unir y separar solicitudes, definir paradas y su orden, editar o anular solicitudes, ver reportes.',
           'Coordinación, Reportes'],
          ['Conductor', 'CHTRANS',
           'Ver únicamente los servicios asignados a su placa, ver la ruta y paradas ordenadas, marcar Iniciar/Finalizar, agregar pasajeros de emergencia si hay cupo.',
           'Mis servicios'],
          ['Administrador', 'ADTRANS',
           'Gestionar catálogos (usuarios, conductores, unidades, puntos, motivos), generar reportes y auditoría. Además puede operar como supervisor, coordinador o conductor.',
           'Solicitudes, Coordinación, Mis servicios, Reportes, Administración'],
      ],
      anchos=[3.2, 2.0, 7.5, 4.0])

d.add_heading('3. Flujo general del proceso', 1)
d.add_paragraph('El ciclo de vida de una solicitud de movilidad es el siguiente:')
tabla(d,
      ['#', 'Actor', 'Acción', 'Estado resultante'],
      [
          ['1', 'Supervisor', 'Crea la solicitud: punto de partida, punto de llegada, fecha, hora, cantidad de personas, motivo y si es emergencia médica.', 'PENDIENTE'],
          ['2', 'Sistema', 'Publica la solicitud en el listado. Las emergencias aparecen primero y en rojo.', 'PENDIENTE'],
          ['3', 'Otro supervisor (opcional)', 'Ve una movilidad programada y se acopla si la ruta pasa por su punto y hay cupo.', 'ASIGNADO'],
          ['4', 'Jefe de Transporte', 'Revisa solicitudes pendientes y la disponibilidad de las combis (badges verde/ámbar/rojo con el cupo).', '—'],
          ['5', 'Jefe de Transporte', 'Asigna placa a cada solicitud: manual, automática (primera unidad con cupo) o múltiple si supera la capacidad de una combi. Puede unir varias solicitudes en un solo traslado y definir paradas intermedias con su orden.', 'ASIGNADO'],
          ['6', 'Sistema', 'Notifica al conductor de la unidad y al supervisor que registró la solicitud (campana y notificación push del navegador).', 'ASIGNADO'],
          ['7', 'Conductor', 'Visualiza el servicio con la ruta y las paradas ordenadas (punto, suben, bajan) y el mapa.', 'ASIGNADO'],
          ['8', 'Conductor', 'Presiona "Iniciar" al comenzar el recorrido.', 'EN_RUTA'],
          ['9', 'Conductor', 'Recoge y deja personal según las paradas. Puede agregar pasajeros de emergencia si la unidad tiene cupo libre.', 'EN_RUTA'],
          ['10', 'Conductor', 'Presiona "Finalizar" al terminar; el sistema guarda la hora de fin.', 'REALIZADO'],
          ['11', 'Sistema', 'Sincroniza con el servidor y registra auditoría (quién creó, modificó, asignó o cambió estados).', 'REALIZADO'],
          ['12', 'Administrador / Jefe', 'Genera reportes de solicitudes, eficiencia de unidades, conductores, áreas y auditoría.', '—'],
      ],
      anchos=[0.8, 3.4, 9.5, 3.0])
d.add_paragraph('Estados de una solicitud: PENDIENTE → ASIGNADO → EN_RUTA → REALIZADO. '
                'En cualquier momento puede pasar a ANULADO (por el supervisor que la creó o por el jefe de transporte).')

d.add_heading('4. Uso por rol', 1)

d.add_heading('4.1 Supervisor', 2)
d.add_paragraph('Pantalla: Solicitudes. Al ingresar ve el listado de movilidades con prioridad, '
                'fecha/hora, puntos, cantidad, placa asignada, cupo del vehículo y estado.')
viñetas(d, [
    'Crear solicitud: botón "Crear" (arriba a la derecha). Completa fecha, hora, punto de partida, '
    'punto de llegada, cantidad de personas, motivo, observación y la casilla "Emergencia médica" si aplica.',
    'Filtrar por fecha: el campo de fecha sobre la tabla; "Ver todas" limpia el filtro.',
    'Ver detalle: clic en la fila. Si la solicitud pertenece a un traslado, se muestran la ruta, '
    'las paradas con cantidades y las demás solicitudes del traslado.',
    'Acoplarse: en el detalle de una solicitud PENDIENTE aparecen los traslados con cupo '
    'compatibles con su ruta (máx. 10 km de distancia a la ruta); basta un clic para unirse.',
    'Editar: ícono de lápiz, disponible mientras no esté realizada ni anulada.',
    'Anular: ícono de X rojo, con confirmación.',
    'Indicador de cupo: la columna CUPO muestra ocupada/capacidad en verde (disponible), '
    'ámbar (casi lleno) o rojo (lleno).',
])

d.add_heading('4.2 Jefe de Transporte (Coordinador)', 2)
d.add_paragraph('Pantalla: Coordinación. A la izquierda, todas las solicitudes; a la derecha, los traslados activos.')
viñetas(d, [
    'Disponibilidad de unidades: badges superiores con placa y cupo libre por color '
    '(verde libre, ámbar parcial, rojo lleno). Los selectores de placa muestran el cupo y bloquean unidades llenas.',
    'Asignar unidad: elegir placa en el selector de la fila y presionar ✓.',
    'Asignación automática (⚡): el sistema toma la primera unidad con cupo suficiente.',
    'Asignación múltiple (icono de dos ventanas): para solicitudes que superan una combi, '
    'reparte la cantidad entre varias unidades.',
    'Unir solicitudes: marcar 2 o más con los checkboxes, elegir la placa del traslado y presionar "Unir". '
    'Se crea un traslado T-N que conserva la identidad de cada solicitud.',
    'Separar: el botón de desenlazar devuelve la solicitud a PENDIENTE sin placa.',
    'Definir paradas: "Editar paradas" en el traslado permite agregar puntos con cantidades '
    'que suben/bajan y ordenarlos con flechas.',
    'Mapa y OSM: cada traslado puede visualizarse en el mapa interno o abrirse en OpenStreetMap con la ruta.',
    'Editar / Anular solicitudes: íconos de lápiz y X en cada fila (gestión total sobre las solicitudes).',
    'Reasignar por urgencia: una solicitud asignada puede cambiarse de placa repitiendo la asignación.',
])

d.add_heading('4.3 Conductor', 2)
d.add_paragraph('Pantalla: Mis servicios. Solo ve los servicios asignados a su placa — nunca el listado general.')
viñetas(d, [
    'La franja superior muestra la "Ruta del día" con la secuencia de puntos.',
    'Iniciar: cambia el servicio a EN_RUTA y registra la hora de inicio.',
    'Finalizar: cambia a REALIZADO y registra la hora de fin.',
    'Agregar pasajeros (icono persona+): registra pasajeros extra de emergencia si hay cupo libre; '
    'la solicitud queda marcada como emergencia y se notifica al supervisor.',
    'Ver mapa: muestra la ruta con las paradas ordenadas (número, punto, suben, bajan), '
    'la ubicación actual del conductor y opción de abrir la navegación en OpenStreetMap.',
    'Columna CUPO: ocupada/capacidad del vehículo con color.',
])

d.add_heading('4.4 Administrador', 2)
d.add_paragraph('Pantallas: Administración y Reportes.')
viñetas(d, [
    'Administración con pestañas: Usuarios (supervisor, coordinador, administrador), Conductores '
    '(con placa asignada), Unidades (placa y capacidad), Puntos (nombre y coordenadas para el mapa) y Motivos.',
    'Cada pestaña permite crear, editar y desactivar registros.',
    'Reportes: rango de fechas, filtros por supervisor/placa/estado, indicadores (solicitudes, personas, '
    'realizados, emergencias, duración y espera promedio, puntualidad), eficiencia por unidad, '
    'servicios por conductor, movimiento por área, exportación a CSV e impresión/PDF.',
    'Auditoría: botón "Cargar auditoría" muestra el historial de cambios con paginación.',
])

d.add_heading('5. Funciones transversales', 1)
viñetas(d, [
    'Notificaciones: la campana del encabezado muestra asignaciones, cambios de estado y pasajeros agregados. '
    'Si el navegador lo permite, también se emiten notificaciones push.',
    'Cambiar clave: ícono de llave junto al nombre del usuario.',
    'Modo offline: si no hay internet, las operaciones se guardan en el equipo y se sincronizan '
    'automáticamente al recuperar la conexión. El badge amarillo del encabezado indica operaciones pendientes. '
    'Las últimas consultas quedan en caché para poder consultar datos sin conexión.',
])

d.add_heading('6. Reglas de negocio', 1)
tabla(d,
      ['Regla', 'Descripción'],
      [
          ['Capacidad', 'Cada combi tiene capacidad (~15). El sistema impide asignar, unir, acoplar o agregar pasajeros por encima del cupo.'],
          ['Prioridad', 'Las emergencias médicas tienen prioridad máxima. Entre solicitudes normales rige "primero en llegar, primero atendido" (orden de registro).'],
          ['Trazabilidad', 'Una solicitud conserva su identidad aunque se una a un traslado (cabecera + detalles). Todo cambio queda en auditoría.'],
          ['Visibilidad', 'El conductor solo ve los servicios de su placa. El supervisor ve el listado general pero solo gestiona sus solicitudes.'],
          ['Compatibilidad de ruta', 'Acoplarse exige cupo y que origen/destino estén a menos de 10 km de la ruta del traslado.'],
          ['Alcance', 'Solicitudes de 50-100+ personas se gestionan con buses externos, fuera del aplicativo.'],
      ],
      anchos=[3.5, 13.0])

d.save(os.path.join(OUT, 'Flujo_y_Uso_por_Rol_Transporte_Personal.docx'))
print('OK doc 1')

# ============================================================
# DOCUMENTO 2: MANUAL DE USUARIO (DETALLADO CON CAPTURAS)
# ============================================================
m = Document()
estilo_base(m)
portada(m, 'Manual de Usuario', 'Aplicativo de Transporte de Personal Interno')

# ---------------- 1. INTRODUCCIÓN ----------------
m.add_heading('1. Introducción', 1)
m.add_paragraph(
    'Este manual explica paso a paso, con capturas reales de pantalla, cómo utilizar el aplicativo '
    'web de transporte de personal. El sistema reemplaza la coordinación por WhatsApp y hojas de '
    'cálculo: un supervisor solicita una movilidad, el jefe de transporte asigna la combi y las '
    'paradas, el conductor ejecuta el viaje y la dirección/administración consulta reportes.')
m.add_paragraph('Requisitos para usar el aplicativo:')
viñetas(m, [
    'Un navegador moderno (Chrome, Edge o similar), en computadora o celular.',
    'Un usuario y contraseña asignados por el administrador del sistema.',
    'Conexión a internet para entrar; si la conexión se pierde mientras trabaja, la aplicación '
    'guarda sus acciones en el equipo y las envía automáticamente al volver (modo offline).',
    'Opcionalmente, aceptar el permiso de notificaciones del navegador para recibir avisos '
    'aunque esté en otra pestaña.',
])
m.add_paragraph(
    'Convenio del manual: cada figura muestra la pantalla real del sistema y debajo se indica su '
    'nombre ("Fig. N — descripción"). Los textos entre comillas corresponden a botones o '
    'etiquetas tal como aparecen en pantalla.')

# ---------------- 2. ACCESO ----------------
m.add_heading('2. Acceso al sistema', 1)
pasos(m, [
    'Abra el navegador e ingrese a la dirección del aplicativo (por ejemplo http://localhost:4200 '
    'o la URL interna que le indique el administrador).',
    'Escriba su Usuario y su Clave en el formulario de ingreso.',
    'Presione el botón "Ingresar".',
    'El sistema lo lleva automáticamente a su pantalla principal según su rol.',
])
imagen(m, '01_login.png', 'Fig. 1 — Pantalla de inicio de sesión. El ojo junto al campo Clave permite ver lo que escribe.')
m.add_paragraph('La pantalla inicial que verá depende de su rol:')
tabla(m,
      ['Rol', 'Pantalla inicial', 'Opciones del menú'],
      [
          ['Supervisor', 'Solicitudes', 'Solicitudes'],
          ['Jefe de Transporte', 'Coordinación', 'Coordinación, Reportes'],
          ['Conductor', 'Mis servicios', 'Mis servicios'],
          ['Administrador', 'Coordinación', 'Solicitudes, Coordinación, Mis servicios, Reportes, Administración'],
      ],
      anchos=[4.0, 4.0, 8.5])
m.add_paragraph(
    'Si el usuario o la clave son incorrectos aparece un aviso rojo bajo el formulario. Verifique '
    'los datos o pida al administrador que revise su cuenta en Administración → Usuarios.')
imagen(m, '00_login_error.png', 'Fig. 2 — Mensaje de error cuando las credenciales no son válidas.')
m.add_paragraph(
    'Para cerrar sesión presione el ícono de salida (puerta con flecha) en la esquina superior '
    'derecha del encabezado, junto a su nombre.')

# ---------------- 3. ELEMENTOS COMUNES ----------------
m.add_heading('3. Elementos comunes de la pantalla', 1)
m.add_paragraph(
    'Toda la aplicación comparte la misma estructura: a la izquierda el menú lateral morado con '
    'las opciones de su rol; arriba el encabezado con el nombre del sistema, su nombre de usuario '
    'con su rol, y los íconos de notificaciones, cambio de clave y salir.')
tabla(m,
      ['Elemento', 'Dónde está', 'Para qué sirve'],
      [
          ['Menú lateral', 'Columna izquierda (morada)', 'Navegar entre los módulos de su rol. El botón ☰ (arriba) lo pliega para ganar espacio.'],
          ['Campana 🔔', 'Encabezado, derecha', 'Abre el panel de notificaciones. Un número rojo indica cuántas no ha leído; toque una para marcarla como leída.'],
          ['Badge amarillo', 'Encabezado, junto a la campana', 'Aparece solo si hay operaciones guardadas sin internet pendientes de sincronizar.'],
          ['Llave 🔑', 'Encabezado, junto a su nombre', 'Abre el formulario para cambiar su contraseña (mínimo 6 caracteres).'],
          ['Nombre + rol', 'Encabezado, derecha', 'Muestra su nombre y la etiqueta de su rol (SUPERVISOR, COORDINADOR, CONDUCTOR o ADMINISTRADOR).'],
          ['Salir ➜', 'Último ícono del encabezado', 'Cierra la sesión y vuelve a la pantalla de ingreso.'],
      ],
      anchos=[3.2, 4.3, 9.0])
m.add_paragraph('El menú lateral muestra solo las opciones de su rol:')
imagen(m, 'menu_supervisor.png', 'Fig. 3 — Menú del Supervisor: solo "Solicitudes".', Cm(5.5))
imagen(m, 'menu_coordinador.png', 'Fig. 4 — Menú del Coordinador: "Coordinación" y "Reportes".', Cm(5.5))
imagen(m, 'menu_conductor.png', 'Fig. 5 — Menú del Conductor: solo "Mis servicios".', Cm(5.5))
imagen(m, 'menu_admin.png', 'Fig. 6 — Menú del Administrador: acceso a todos los módulos.', Cm(5.5))
imagen(m, '05_notificaciones.png',
       'Fig. 7 — Panel de notificaciones al presionar la campana. Avisos de asignaciones, uniones, '
       'cambios de estado y pasajeros agregados; el contador rojo indica las no leídas.')
imagen(m, '15_cambiar_clave.png',
       'Fig. 8 — El ícono de llave (arriba a la derecha) despliega el formulario de cambio de '
       'contraseña: clave actual, clave nueva y confirmación.')
m.add_heading('3.1 Trabajar sin internet (modo offline)', 2)
m.add_paragraph(
    'Si pierde la conexión, la aplicación sigue funcionando con la última información cargada: '
    'puede consultar listados que ya haya visto y ejecutar acciones (crear, asignar, cambiar '
    'estado, etc.). Estas acciones se guardan en el equipo y aparecen contadas en el badge '
    'amarillo del encabezado. Al recuperar internet se envían solas en el mismo orden y el badge '
    'desaparece cuando todo quedó sincronizado.')

# ---------------- 4. SOLICITUDES ----------------
m.add_heading('4. Módulo Solicitudes (Supervisor)', 1)
m.add_paragraph(
    'Es la pantalla principal del supervisor. Muestra el listado de movilidades programadas con '
    'todas sus columnas y permite crear, editar, anular y acoplar solicitudes.')
imagen(m, '02_solicitudes.png',
       'Fig. 9 — Listado de solicitudes. Arriba: filtro por fecha y botón "Crear". La fila en '
       'rojo correspondería a una emergencia; las agrupadas muestran su traslado.')

m.add_heading('4.1 Entender el listado', 2)
tabla(m,
      ['Columna', 'Qué muestra'],
      [
          ['PRIORIDAD', 'NORMAL o EMERGENCIA. Las emergencias van primero en la lista y la fila se pinta en rojo.'],
          ['NOMBRE', 'Quién registró la solicitud.'],
          ['FECHA / HORA', 'Fecha y hora programada del viaje.'],
          ['ÁREA', 'Área del solicitante.'],
          ['P. PARTIDA / P. LLEGADA', 'Punto de origen y punto de destino (del catálogo de Puntos).'],
          ['CANTIDAD', 'Número de personas a trasladar.'],
          ['MOTIVO', 'Motivo del viaje (capacitación, actividad, salud, etc.).'],
          ['PLACA', 'Unidad asignada (vacío mientras esté pendiente).'],
          ['CUPO', 'Ocupadas/capacidad de la unidad asignada. Verde = hay espacio, ámbar = casi llena, rojo = llena.'],
          ['ESTADO', 'PENDIENTE, ASIGNADO, EN_RUTA, REALIZADO o ANULADO.'],
          ['Acciones', 'Lápiz = editar la solicitud. X roja = anularla (pide confirmación).'],
      ],
      anchos=[4.0, 12.5])
m.add_paragraph(
    'Sobre la tabla hay un campo "Fecha": al elegir un día la lista se limita a las movilidades '
    'de esa fecha, y aparece el botón "Ver todas" para quitar el filtro. Clic sobre cualquier '
    'fila abre el detalle de la solicitud.')
imagen(m, '02b_filtro_fecha.png', 'Fig. 10 — Listado filtrado por fecha con la opción "Ver todas" para limpiar el filtro.')

m.add_heading('4.2 Crear una solicitud', 2)
pasos(m, [
    'Presione el botón "Crear" en la parte superior derecha.',
    'Su nombre aparece ya cargado en el campo NOMBRE (no se puede modificar).',
    'Complete los campos obligatorios (marcados con *):',
])
tabla(m,
      ['Campo', 'Obligatorio', 'Descripción'],
      [
          ['FECHA PROGRAMADA', 'Sí', 'Día en que se realizará el viaje.'],
          ['Hr. PROGRAMADA', 'Sí', 'Hora de salida.'],
          ['P. PARTIDA', 'Sí', 'Punto donde sube el personal (lista del catálogo de Puntos).'],
          ['P. LLEGADA', 'Sí', 'Punto de destino.'],
          ['CANTIDAD', 'Sí', 'Personas a trasladar (mínimo 1; con flechas para subir/bajar).'],
          ['MOTIVO', 'Sí', 'Motivo del viaje (lista del catálogo de Motivos).'],
          ['OBSERVACIÓN', 'No', 'Comentario libre para el jefe de transporte o el conductor.'],
          ['Emergencia médica', 'No', 'Marcar solo para urgencias: le da prioridad máxima.'],
      ],
      anchos=[4.5, 2.5, 9.5])
pasos(m, [
    'Presione "Save". El botón "Cancel" cierra sin guardar.',
    'La solicitud aparece en el listado con estado PENDIENTE; el jefe de transporte ya puede verla.',
])
imagen(m, '03_crear_solicitud.png',
       'Fig. 11 — Formulario de nueva solicitud. Todos los campos con * deben llenarse para habilitar "Save".')

m.add_heading('4.3 Ver el detalle de una solicitud', 2)
m.add_paragraph(
    'Haga clic sobre la fila de la solicitud. Se abre una ventana con todos sus datos: prioridad, '
    'solicitante, fecha y hora, área, puntos, cantidad, motivo, observación, placa, estado y si ya '
    'fue realizada. Si la solicitud está unida a un traslado, la ventana muestra además la ruta, '
    'las paradas con cuántos suben y bajan en cada una, y las demás solicitudes que comparten el viaje.')
imagen(m, '04_detalle_solicitud.png', 'Fig. 12 — Detalle completo de una solicitud (clic sobre la fila para abrirlo).')

m.add_heading('4.4 Acoplarse a una movilidad existente', 2)
m.add_paragraph(
    'Si otro supervisor ya tiene un traslado programado que pasa cerca de su punto de partida y '
    'destino, puede unir su solicitud a ese viaje sin esperar una nueva asignación.')
pasos(m, [
    'Cree su solicitud y ábrala haciendo clic en su fila.',
    'Al final del detalle aparece la sección "Acoplar a traslado con cupo" con los traslados '
    'compatibles (mismo corredor: origen y destino a menos de 10 km de la ruta).',
    'Cada botón muestra el traslado (T-N), la placa y su ruta. Presione el que le sirva.',
    'El sistema verifica que haya cupo; si lo hay, su solicitud queda ASIGNADA a ese traslado '
    'conservando su registro individual.',
])
m.add_paragraph(
    'Si no aparece la sección "Acoplar", es porque no hay traslados con cupo compatibles con su '
    'ruta; su solicitud esperará la asignación normal del jefe de transporte.')
imagen(m, '04b_acoplar.png',
       'Fig. 13 — Detalle de una solicitud PENDIENTE con el botón del traslado compatible (T-2 · TP-001 · ruta) listo para acoplarse.')

m.add_heading('4.5 Editar o anular una solicitud', 2)
viñetas(m, [
    'Editar (lápiz): abre el mismo formulario de creación con los datos cargados para corregirlos. '
    'Disponible mientras la solicitud no esté realizada ni anulada.',
    'Anular (X roja): el sistema pide confirmación antes de cancelar la solicitud. La acción no '
    'se puede deshacer y queda registrada en la auditoría.',
])
imagen(m, '03b_editar_solicitud.png', 'Fig. 14 — Edición de una solicitud: el formulario se abre con los datos actuales.')
imagen(m, '03c_anular_confirm.png', 'Fig. 15 — Confirmación antes de anular una solicitud.')

m.add_heading('4.6 Estados de una solicitud', 2)
tabla(m,
      ['Estado', 'Qué significa', 'Quién lo produce'],
      [
          ['PENDIENTE', 'Recién creada, aún sin unidad asignada.', 'Supervisor'],
          ['ASIGNADO', 'Ya tiene placa (solita o dentro de un traslado).', 'Jefe de transporte / acople'],
          ['EN_RUTA', 'El conductor inició el recorrido.', 'Conductor'],
          ['REALIZADO', 'El conductor finalizó el viaje; se guardan las horas.', 'Conductor'],
          ['ANULADO', 'Solicitud cancelada.', 'Supervisor / jefe de transporte'],
      ],
      anchos=[3.0, 9.0, 4.5])

# ---------------- 5. COORDINACIÓN ----------------
m.add_heading('5. Módulo Coordinación (Jefe de Transporte)', 1)
m.add_paragraph(
    'Es el centro de operaciones del día. La pantalla se divide en dos paneles: a la izquierda '
    'todas las solicitudes con sus acciones, y a la derecha los traslados activos con su ocupación.')
imagen(m, '06_coordinacion.png',
       'Fig. 16 — Coordinación: solicitudes a la izquierda, "Traslados activos" a la derecha y '
       'badges de disponibilidad de unidades sobre la tabla.')

m.add_heading('5.1 Ver la disponibilidad de las unidades', 2)
m.add_paragraph(
    'Sobre la tabla hay un badge por cada combi con la forma "PLACA · libres/capacidad", por '
    'ejemplo "TP-001 · 15/15 libres". El color resume el estado:')
viñetas(m, [
    'Verde: la unidad está libre o con bastante cupo.',
    'Ámbar: parcialmente ocupada.',
    'Rojo: llena — no acepta más pasajeros.',
])
m.add_paragraph(
    'Los selectores de placa en cada fila muestran el mismo cupo ("TP-002 · 15/15 libres") y '
    'bloquean automáticamente las unidades sin espacio, para que no pueda asignar más personas '
    'de las que caben.')

m.add_heading('5.2 Asignar unidad a una solicitud', 2)
m.add_paragraph('Cada fila de solicitud sin traslado tiene estos botones de acción:')
tabla(m,
      ['Botón', 'Acción'],
      [
          ['Selector "Placa" + ✓', 'Elige la unidad en la lista (muestra el cupo de cada una) y presione el check para asignarla.'],
          ['⚡ (rayo)', 'Asignación automática: el sistema toma la primera unidad con cupo suficiente.'],
          ['⧉ (ventanas)', 'Asignación múltiple: reparte la cantidad entre varias unidades cuando una sola no alcanza.'],
          ['Lápiz', 'Editar los datos de la solicitud.'],
          ['X', 'Anular la solicitud.'],
      ],
      anchos=[5.0, 11.5])
imagen(m, '06b_asignacion_placa.png',
       'Fig. 17 — Fila de solicitud con la placa seleccionada en el selector (muestra el cupo) '
       'y los botones de asignar ✓, automática ⚡ y múltiple ⧉.')
pasos(m, [
    'En la columna ACCIONES de la solicitud elija la placa en el selector.',
    'Presione el botón ✓. Confirme el mensaje.',
    'La solicitud pasa a ASIGNADO y el conductor de esa placa recibe una notificación.',
    'Para reasignar a otra unidad (por ejemplo por una urgencia), repita el proceso con otra placa.',
])
m.add_paragraph(
    'Asignación múltiple (⧉): para solicitudes que superan la capacidad de una combi. Se abre un '
    'panel donde agrega tantas unidades como necesite con la cantidad de personas que llevará cada '
    'una; el total repartido debe coincidir exactamente con lo solicitado.')
imagen(m, '06f_asignacion_multiple.png',
       'Fig. 18 — Panel "Distribuir N personas": cada fila es una unidad con la cantidad que le corresponde.')

m.add_heading('5.3 Unir solicitudes en un traslado', 2)
pasos(m, [
    'Marque con los checkboxes de la primera columna dos o más solicitudes compatibles.',
    'Arriba de la tabla aparecen el selector "Placa del traslado" y el botón "Unir (N)".',
    'Elija la placa y presione "Unir". Se crea el traslado T-N que agrupa las solicitudes: cada '
    'una conserva su número, solicitante y cantidad para los reportes.',
])
imagen(m, '06e_traslado_unido.png',
       'Fig. 19 — Traslado T-2 creado: tag de ocupación 10/15, ruta combinada, solicitudes que lo '
       'integran y acciones Editar paradas / Mapa / OSM. Las filas unidas muestran su tag "T-2".')
m.add_paragraph(
    'Separar: en la fila de una solicitud que está dentro de un traslado aparece el botón de '
    'desenlazar (⛓). Tras confirmar, la solicitud vuelve a estado PENDIENTE sin placa y el '
    'traslado recalcula su ocupación.')
imagen(m, '06i_separar_confirm.png', 'Fig. 20 — Confirmación al separar una solicitud de su traslado.')

m.add_heading('5.4 Definir las paradas y la ruta', 2)
pasos(m, [
    'En la tarjeta del traslado (panel derecho) presione "Editar paradas".',
    'Con "Agregar parada" añada filas: en cada una elija el punto (lista con buscador) y escriba '
    'cuántas personas suben y cuántas bajan ahí.',
    'Use las flechas ↑ ↓ para ordenar la secuencia de recogida y el basurero para quitar una parada.',
    'Presione "Guardar ruta". El conductor verá las paradas en ese orden en su aplicación.',
])
imagen(m, '06d_editor_paradas.png',
       'Fig. 21 — Editor de paradas del traslado: punto, suben, bajan, botones de orden y quitar.')
m.add_paragraph(
    'Con "Mapa" se despliega el recorrido dibujado sobre OpenStreetMap con los puntos de la '
    'ruta; "OSM" abre la misma ruta en una pestaña nueva para compartirla o navegarla.')
imagen(m, '06h_mapa_traslado.png',
       'Fig. 22 — Mapa del traslado con los puntos marcados y la tarjeta del traslado con su ruta y paradas.')

m.add_heading('5.5 Editar o anular solicitudes', 2)
m.add_paragraph(
    'A diferencia del supervisor, el jefe de transporte puede modificar cualquier solicitud. El '
    'lápiz abre un panel con todos los campos (fecha, hora, puntos, cantidad, motivo, observación '
    'y la marca de emergencia) para corregirlos; la X la anula con confirmación. Todo queda en auditoría.')
imagen(m, '06g_editar_solicitud.png', 'Fig. 23 — Panel de edición de solicitud del coordinador.')

# ---------------- 6. MIS SERVICIOS ----------------
m.add_heading('6. Módulo Mis servicios (Conductor)', 1)
m.add_paragraph(
    'El conductor solo ve los servicios asignados a la placa de su unidad (mostrada bajo el título, '
    'por ejemplo "Unidad: TP-002"), ordenados con las emergencias primero. El botón "Actualizar" '
    'recarga la lista. La franja azul superior resume la ruta del día con la secuencia de puntos.')
imagen(m, '07_mis_servicios.png',
       'Fig. 24 — Mis servicios: solo los viajes de la placa del conductor, con cupo, prioridad, '
       'estado y botones de acción por fila.')
tabla(m,
      ['Elemento', 'Qué hace'],
      [
          ['Iniciar (verde)', 'Aparece en servicios ASIGNADO. Lo presiona al arrancar: pasa a EN_RUTA y guarda la hora de inicio.'],
          ['Finalizar (verde)', 'Aparece en servicios EN_RUTA. Lo presiona al terminar: pasa a REALIZADO y guarda la hora de fin.'],
          ['Persona+ (rojo)', 'Agregar pasajeros de emergencia si hay cupo libre.'],
          ['Mapa (celeste)', 'Ver la ruta del servicio con sus paradas ordenadas.'],
          ['CUPO', 'Ocupadas/capacidad del vehículo en ese servicio, con color.'],
      ],
      anchos=[4.0, 12.5])
pasos(m, [
    'Al llegar a su turno, abra "Mis servicios" y revise los viajes asignados: origen, destino, '
    'cantidad de personas y motivo.',
    'Antes de arrancar, presione el ícono del mapa para conocer la ruta completa.',
    'Al iniciar el recorrido presione "Iniciar". El estado cambia a EN_RUTA.',
    'Recorra las paradas en el orden indicado: la tabla muestra cuántos suben y cuántos bajan en cada punto.',
    'Si en el camino debe subir pasajeros de emergencia, use el botón persona+ (ver abajo).',
    'Al terminar el viaje presione "Finalizar". El servicio queda REALIZADO.',
    'Sin señal puede seguir operando: los cambios se guardan en el equipo y se envían al recuperar '
    'internet (badge amarillo en el encabezado).',
])
imagen(m, '08_ruta_paradas.png',
       'Fig. 25 — Ruta del servicio: mapa con la trayectoria y tabla de paradas ordenadas '
       '(# / punto / suben / bajan). "Abrir en OSM" lleva la ruta a la navegación.')
m.add_heading('6.1 Agregar pasajeros de emergencia', 2)
pasos(m, [
    'En un servicio ASIGNADO o EN_RUTA presione el botón rojo de persona+.',
    'El panel muestra el cupo libre de la unidad ("Cupo libre: X de 15").',
    'Escriba cuántos pasajeros adicionales suben y presione "Agregar".',
    'Si la cantidad supera el cupo disponible el sistema lo rechaza. Al guardar, la solicitud '
    'queda marcada como emergencia y el supervisor recibe una notificación.',
])
imagen(m, '09_agregar_pasajeros.png', 'Fig. 26 — Panel de pasajeros de emergencia con el cupo libre visible.')

# ---------------- 7. REPORTES ----------------
m.add_heading('7. Módulo Reportes (Jefe / Administrador)', 1)
m.add_paragraph('Genera el histórico de movilidades y los indicadores de operación.')
pasos(m, [
    'Elija el rango "Desde" / "Hasta" y presione "Generar".',
    'Opcionalmente refine con los filtros Supervisor, Placa y Estado (se aplican al instante '
    'sobre la tabla, los totales y el CSV).',
    'Revise las tarjetas de indicadores y las tablas de eficiencia.',
    'Use "Exportar CSV" para descargar el detalle (compatible con Excel) o "Imprimir / PDF" '
    'para la versión imprimible.',
])
imagen(m, '10_reportes.png',
       'Fig. 27 — Reportes: filtros de fecha y supervisor/placa/estado, totales rápidos y tarjetas '
       'de indicadores de operación.')
tabla(m,
      ['Indicador', 'Qué mide'],
      [
          ['Solicitudes / Personas', 'Total de movilidades y de personas trasladadas en el rango.'],
          ['Realizados / Emergencias', 'Viajes completados y cuántos fueron de emergencia.'],
          ['Duración promedio', 'Tiempo medio entre Iniciar y Finalizar.'],
          ['Espera promedio', 'Tiempo medio entre la creación de la solicitud y su atención.'],
          ['Puntualidad ±10 min', 'Porcentaje de servicios realizados dentro de 10 minutos de la hora programada.'],
          ['Eficiencia por unidad', 'Viajes, personas y porcentaje de ocupación por combi.'],
          ['Servicios por conductor', 'Viajes y personas atendidas por cada conductor.'],
          ['Movimiento por área', 'Solicitudes y personas agrupadas por el área que las pidió.'],
      ],
      anchos=[5.0, 11.5])
imagen(m, '10b_reporte_tabla.png', 'Fig. 28 — Detalle de solicitudes del reporte con placa, prioridad y realización por fila.')
m.add_heading('7.1 Auditoría', 2)
m.add_paragraph(
    'El botón "Cargar auditoría" muestra el historial de todas las acciones del sistema: inicios '
    'de sesión, creación y edición de solicitudes, asignaciones, uniones, separaciones, cambios de '
    'estado y pasajeros agregados, con fecha, entidad afectada y usuario. La tabla pagina de a 5 '
    'filas y permite cambiar a 10 o 20 por página.')
imagen(m, '11_auditoria.png', 'Fig. 29 — Auditoría del sistema con su paginador (5 filas por página).')

# ---------------- 8. ADMINISTRACIÓN ----------------
m.add_heading('8. Módulo Administración (Administrador)', 1)
m.add_paragraph(
    'Gestiona los catálogos del sistema en cinco pestañas. En todas funciona igual: el formulario '
    'de la parte superior sirve para crear (botón "Crear") y para editar: al presionar "Editar" en '
    'una fila, sus datos se cargan en el formulario y el botón cambia a "Actualizar". "Desactivar" '
    'inhabilita el registro sin borrarlo; en usuarios, si deja la clave vacía al editar se conserva la actual.')
tabla(m,
      ['Pestaña', 'Campos del formulario', 'Para qué sirve'],
      [
          ['Usuarios', 'Usuario, nombre, rol (supervisor/coordinador/administrador), área y clave.',
           'Cuentas del personal que solicita, coordina o administra.'],
          ['Conductores', 'Usuario, nombre, placa de la unidad asignada, área y clave.',
           'Cuentas de conductor; la placa define qué servicios verá en "Mis servicios".'],
          ['Unidades', 'Placa y capacidad (≈15).',
           'Las combis disponibles para asignar.'],
          ['Puntos', 'Nombre, latitud y longitud.',
           'Orígenes/destinos y paradas. Las coordenadas alimentan el mapa y la validación de rutas al acoplarse.'],
          ['Motivos', 'Nombre.',
           'Catálogo de motivos de viaje que elige el supervisor.'],
      ],
      anchos=[2.8, 6.7, 7.0])
imagen(m, '12_administracion_usuarios.png', 'Fig. 30 — Pestaña Usuarios: formulario arriba y tabla con Editar/Desactivar por fila.')
imagen(m, '13_administracion_conductores.png', 'Fig. 31 — Pestaña Conductores: cada conductor tiene su placa asignada.')
imagen(m, '17_unidades.png', 'Fig. 32 — Pestaña Unidades: placa y capacidad de cada combi.')
imagen(m, '14_administracion_puntos.png', 'Fig. 33 — Pestaña Puntos: nombre, latitud y longitud de cada punto.')
imagen(m, '18_motivos.png', 'Fig. 34 — Pestaña Motivos del catálogo.')

# ---------------- 9. GUÍA RÁPIDA ----------------
m.add_heading('9. Guía rápida por rol', 1)
tabla(m,
      ['Rol', 'Flujo típico'],
      [
          ['Supervisor',
           '1. Ingresar → Solicitudes. 2. "Crear" y llenar el formulario. 3. (Opcional) Abrir el detalle y '
           'acoplarse a un traslado con cupo. 4. Seguir el estado en la columna ESTADO. 5. Editar o anular si cambia el plan.'],
          ['Jefe de Transporte',
           '1. Ingresar → Coordinación. 2. Revisar badges de cupo por unidad. 3. Asignar placa (✓, ⚡ o ⧉) '
           'o unir solicitudes compatibles en un traslado. 4. "Editar paradas" para ordenar el recorrido. '
           '5. Atender urgencias reasignando. 6. Revisar Reportes al cierre.'],
          ['Conductor',
           '1. Ingresar → Mis servicios. 2. Revisar la ruta y las paradas con el ícono de mapa. '
           '3. "Iniciar" al arrancar. 4. Persona+ solo para pasajeros de emergencia con cupo. '
           '5. "Finalizar" al terminar.'],
          ['Administrador',
           '1. Mantener catálogos en Administración (usuarios, conductores con su placa, unidades, '
           'puntos con coordenadas, motivos). 2. Revisar Reportes y Auditoría. 3. Puede operar '
           'además cualquier pantalla de los demás roles.'],
      ],
      anchos=[3.5, 13.0])

# ---------------- 10. GLOSARIO ----------------
m.add_heading('10. Glosario', 1)
tabla(m,
      ['Término', 'Significado'],
      [
          ['Solicitud', 'Pedido de transporte creado por un supervisor.'],
          ['Traslado (T-N)', 'Agrupación de una o más solicitudes bajo una misma placa, con su ruta y paradas ordenadas.'],
          ['Acoplarse', 'Unir una solicitud nueva a un traslado ya programado que tenga cupo y ruta compatible (máx. 10 km).'],
          ['Parada', 'Punto intermedio de la ruta donde sube o baja personal.'],
          ['Cupo', 'Asientos libres de la unidad (ocupadas/capacidad).'],
          ['Emergencia', 'Solicitud de máxima prioridad: se muestra en rojo y primero en las listas.'],
          ['Badge', 'Etiqueta de color que resume un estado (verde/ámbar/rojo).'],
          ['Auditoría', 'Registro de quién hizo cada cambio y cuándo.'],
          ['Modo offline', 'Trabajo sin internet: las acciones se guardan en el equipo y se sincronizan al volver la conexión.'],
      ],
      anchos=[4.0, 12.5])

# ---------------- 11. PROBLEMAS FRECUENTES ----------------
m.add_heading('11. Problemas frecuentes', 1)
tabla(m,
      ['Situación', 'Qué hacer'],
      [
          ['No puedo ingresar',
           'Verifique usuario y clave (respete mayúsculas). Si persiste, pida al administrador revisar '
           'su cuenta en Administración → Usuarios.'],
          ['No veo el módulo que necesito',
           'El menú depende del rol: confirme con el administrador que su usuario tiene el rol correcto.'],
          ['El conductor no ve servicios',
           'Verifique que el usuario tenga placa asignada (Administración → Conductores) y que el jefe '
           'haya asignado la solicitud a esa placa.'],
          ['"No hay cupo suficiente"',
           'La unidad está llena: use otra placa, la asignación múltiple (⧉) o libere espacio separando '
           'otra solicitud.'],
          ['El selector de placa no deja elegir una unidad',
           'Esa unidad está llena (aparece deshabilitada con su cupo). Elija otra o use asignación múltiple.'],
          ['Trabajé sin internet',
           'Los cambios se guardaron en el equipo; al volver la conexión se envían solos y el badge '
           'amarillo del encabezado desaparece.'],
          ['No me deja acoplarme a un traslado',
           'El traslado no tiene cupo o su origen/destino está a más de 10 km de la ruta. Espere la '
           'asignación del jefe.'],
          ['El mapa no muestra la ruta',
           'Los puntos necesitan latitud y longitud en Administración → Puntos.'],
          ['No llegan las notificaciones push',
           'Acepte el permiso de notificaciones que pide el navegador; los avisos siempre llegan a la campana.'],
      ],
      anchos=[5.5, 11.0])

m.save(os.path.join(OUT, 'Manual_de_Usuario_Transporte_Personal.docx'))
print('OK doc 2')
